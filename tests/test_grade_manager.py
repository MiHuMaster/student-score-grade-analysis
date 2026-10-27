"""年级版业务验收：只使用合成学生和临时数据库。"""
from datetime import date
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from app import create_app
from app.db import get_db
from app.services.analytics import Scope, _derived_subject_lines, benchmark_metrics, class_metrics, grade_statistics, student_detail
from app.services.importer import commit_batch, delete_import_batch, parse_workbook, preview_workbook
from app.services.standard_exports import export_class_zip


class GradeManagerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.app = create_app({"TESTING": True, "DATABASE": str(self.root / "test.db"),
                               "UPLOAD_FOLDER": str(self.root / "uploads"), "EXPORT_FOLDER": str(self.root / "exports")})
        self.context = self.app.app_context()
        self.context.push()
        self.db = get_db()

    def tearDown(self):
        self.context.pop()
        self.temp.cleanup()

    def workbook(self, records, filename="scores.xlsx", exam="合成月考", school="合成学校"):
        """Record: no, class, Chinese score; optional track, raw/scaled chemistry, input rank."""
        workbook = Workbook()
        schools = workbook.active
        schools.title = "学校信息"
        schools.append(["省份", "学校", "年级", "班级", "方向", "班级人数", "特控目标"])
        students = workbook.create_sheet("学生信息")
        students.append(["学校", "年级", "班级", "学生姓名", "学号", "选科标签"])
        for track in ("物理", "历史"):
            sheet = workbook.create_sheet(track + "方向学生成绩")
            sheet.append(["年级", "考试名称", "考试时间", "学号", "班级", "姓名", "语文", "语文排名",
                          "数学", "英语", track, "化学", "化学赋分", "生物", "生物赋分", "政治", "政治赋分", "地理", "地理赋分"])
        class_keys = set()
        for record in records:
            number, class_name, chinese = record[:3]
            opts = record[3] if len(record) > 3 else {}
            track = opts.get("track", "物理")
            if (class_name, track) not in class_keys:
                schools.append(["合成省", school, "高2027届", class_name, track, 30, 2])
                class_keys.add((class_name, track))
            name = "合成学生" + number
            students.append([school, "高2027届", class_name, name, number, "物化生" if track == "物理" else "史化生"])
            workbook[track + "方向学生成绩"].append([
                "高2027届", exam, date(2026, 10, 1), number, class_name, name, chinese, opts.get("rank"),
                100, 100, 80, opts.get("raw", 60), opts.get("scaled", 80), 60, 80, None, None, None, None,
            ])
        path = self.root / filename
        workbook.save(path)
        workbook.close()
        return path

    def commit(self, path, mode="append"):
        preview = preview_workbook(self.db, path, path.name)
        self.assertEqual(preview["error_count"], 0, preview["issues"])
        result = commit_batch(self.db, preview["batch_key"], mode, warnings_confirmed=True)
        return preview, result

    def test_weighted_means_ties_zero_subject_denominators_and_scope(self):
        path = self.workbook([("001", "1", 0), ("002", "1", 120), ("003", "1", 120), ("004", "2", 100),
                              ("005", "3", 150, {"track": "历史"})])
        _, result = self.commit(path)
        key = result["exam_keys"][0]
        stats = grade_statistics(self.db, key, "物理")
        self.assertEqual(stats["reference_count"], 4)
        self.assertEqual(stats["class_count"], 2)
        self.assertEqual(stats["averages"]["__TOTAL__"], 525)
        self.assertEqual(stats["averages"]["chinese"], 85)
        self.assertEqual(stats["counts"]["chinese"], 4)
        self.assertNotIn("politics", stats["averages"])
        self.assertEqual(sorted(stats["ranks"]["__TOTAL__"].values()), [1, 1, 3, 4])
        classes = class_metrics(self.db, Scope(key, "物理"), include_history=False)["classes"]
        for item in classes:
            filtered = class_metrics(self.db, Scope(key, "物理", class_id=item["class_id"]), include_history=False)["classes"][0]
            self.assertEqual(filtered["benchmark_average"], 525)
            self.assertAlmostEqual(filtered["benchmark_ratio"], item["average"] / 525)
        self.assertEqual(student_detail(self.db, "005")["history"][0]["grade_rank"], 1)
        self.assertEqual(student_detail(self.db, "001")["history"][0]["grade_rank"], 4)

    def test_missing_scaled_blocks_row_and_zero_scaled_is_valid(self):
        path = self.workbook([("001", "1", 100, {"scaled": None}), ("002", "1", 100, {"scaled": 0})])
        payload = parse_workbook(path, self.db)
        missing = [issue for issue in payload.issues if issue.issue_code == "selected_scaled_score_missing"]
        self.assertEqual([issue.row_number for issue in missing], [2])
        preview = preview_workbook(self.db, path, path.name)
        with self.assertRaises(ValueError):
            commit_batch(self.db, preview["batch_key"], "append")
        result = commit_batch(self.db, preview["batch_key"], "append", allow_partial=True, warnings_confirmed=True)
        self.assertEqual(result["committed_rows"], 1)
        stats = grade_statistics(self.db, result["exam_keys"][0], "物理")
        self.assertEqual(stats["averages"]["chemistry"], 0)
        self.assertEqual(stats["counts"]["chemistry"], 1)

    def test_append_rank_warning_uses_existing_population_and_deduplicates(self):
        first = self.workbook([("001", "1", 120), ("002", "1", 120)], "first.xlsx")
        _, result = self.commit(first)
        second = self.workbook([("003", "2", 100, {"rank": 1})], "second.xlsx")
        payload = parse_workbook(second, self.db)
        self.assertTrue(any(issue.issue_code == "grade_rank_mismatch" and "计算 3" in issue.message for issue in payload.issues))
        self.commit(second)
        self.commit(first)
        stats = grade_statistics(self.db, result["exam_keys"][0], "物理")
        self.assertEqual(stats["reference_count"], 3)
        self.assertEqual(student_detail(self.db, "003")["history"][0]["grade_rank"], 3)

    def test_replace_delete_and_exam_history_recalculate(self):
        first = self.workbook([("001", "1", 120), ("002", "1", 100)], "first.xlsx")
        _, old = self.commit(first)
        history = self.workbook([("001", "1", 0), ("002", "1", 100)], "later.xlsx", exam="合成期中")
        _, later = self.commit(history)
        detail = student_detail(self.db, "001")
        self.assertEqual([item["grade_rank"] for item in detail["history"]], [1, 2])
        replacement = self.workbook([("001", "1", 90)], "replace.xlsx", exam="合成期中")
        preview, _ = self.commit(replacement, "replace")
        self.assertEqual(grade_statistics(self.db, later["exam_keys"][0], "物理")["reference_count"], 1)
        self.assertEqual(grade_statistics(self.db, old["exam_keys"][0], "物理")["averages"]["__TOTAL__"], 550)
        delete_import_batch(self.db, preview["batch_key"])
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM exam_results er JOIN exams e ON e.id=er.exam_id WHERE e.exam_key=?", (later["exam_keys"][0],)).fetchone()[0], 0)

    def test_single_school_cannot_be_bypassed_by_partial_import_or_manual_entry(self):
        self.commit(self.workbook([("001", "1", 120)]))
        other = self.workbook([("002", "1", 100)], "other.xlsx", school="另一所合成学校")
        preview = preview_workbook(self.db, other, other.name)
        self.assertGreater(preview["error_count"], 0)
        with self.assertRaises(ValueError):
            commit_batch(self.db, preview["batch_key"], "append", allow_partial=True)
        client = self.app.test_client()
        response = client.post("/api/base-information/classes", json={"school_name":"另一校", "grade_name":"高2027届", "class_name":"2"})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM organizations").fetchone()[0], 1)

    def test_derived_total_preserves_subject_cutoffs_and_replace_rank_warning(self):
        self.commit(self.workbook([("001", "1", 120)]))
        other = self.workbook([("002", "1", 100, {"rank": 2})], "replace.xlsx")
        preview = preview_workbook(self.db, other, other.name)
        self.assertFalse(any(issue["issue_code"] == "grade_rank_mismatch" for issue in preview["issues"]))
        result = commit_batch(self.db, preview["batch_key"], "replace", warnings_confirmed=True)
        self.assertEqual(result["rank_warning_count"], 1)
        key = result["exam_keys"][0]
        response = self.app.test_client().put("/api/settings/score-lines", json={"exam_key":key, "track":"物理", "特控":500})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(_derived_subject_lines(self.db, key, "物理")["chinese"]["特控"], 92.5)

    def test_removed_routes_and_template_fields(self):
        client = self.app.test_client()
        for endpoint in ("/api/export/onsite", "/api/export/summary-no-school", "/api/export/school-reports"):
            self.assertEqual(client.get(endpoint).status_code, 404)
        self.assertEqual(client.put("/api/settings/benchmark", json={}).status_code, 404)
        root = Path(__file__).resolve().parents[1]
        for path in (root / "templates/upload").glob("*.xlsx"):
            workbook = load_workbook(path, read_only=True)
            self.assertNotIn("其他学生数据", workbook.sheetnames)
            self.assertNotIn("额外参考均分", workbook.sheetnames)
            for sheet in workbook:
                for row in sheet:
                    for cell in row:
                        if isinstance(cell.value, str):
                            self.assertNotIn("外部参考", cell.value)
                            self.assertNotIn("外部机构", cell.value)
                if "成绩" in sheet.title:
                    self.assertFalse(any(str(cell.value).isascii() and str(cell.value).isalpha() for cell in sheet[1]))
            workbook.close()

    def test_class_exports_keep_subject_columns_and_grade_ranks(self):
        _, result = self.commit(self.workbook([("001", "1", 120), ("002", "2", 100)]))
        key = result["exam_keys"][0]
        folder = self.root / "exports"
        with ZipFile(export_class_zip(self.db, key, folder)) as archive:
            self.assertEqual(len(archive.namelist()), 2)
            for name in archive.namelist():
                workbook = load_workbook(BytesIO(archive.read(name)), data_only=True)
                sheet = workbook.active
                self.assertEqual(sheet.max_column, 35)
                headers = [cell.value for cell in sheet[1]]
                rank = 1 if sheet.cell(2, headers.index("语文") + 1).value == 120 else 2
                self.assertEqual(sheet.cell(2, headers.index("语文排名") + 1).value, rank)
                self.assertEqual(sheet.cell(2, headers.index("总分赋分排名") + 1).value, rank)
                workbook.close()
