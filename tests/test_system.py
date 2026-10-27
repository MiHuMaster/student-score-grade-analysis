from __future__ import annotations

import tempfile
import unittest
import zipfile
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.units import EMU_to_cm

from app import create_app
from app.config import CORE_RULE_VERSION
from app.db import get_db, run_priority1_data_migration
from app.services.analytics import (
    Scope,
    _derived_subject_lines,
    _line_map,
    _mround_half,
    benchmark_metrics,
    class_metrics,
    comparison_exams,
    filter_options,
    list_exams,
    overview,
    score_bands,
    student_group_matrix,
    student_groups,
    student_detail,
    student_list,
    subject_metrics,
)
from app.services.demo import clear_demo_data, seed_demo_data
from app.services.importer import (
    build_reference_workbook_from_base_information,
    clear_import_cache,
    commit_batch,
    delete_import_batch,
    get_batch_preview,
    list_batch_issues,
    list_batches,
    parse_workbook,
    parse_base_information_workbook,
    prepare_import_workbook,
    preview_workbook,
    save_base_information,
    saved_base_reference_data,
)
from app.services.standard_exports import export_class_zip


class SystemTestCase(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        root = Path(self.tempdir.name)
        self.app = create_app(
            {
                "TESTING": True,
                "DATABASE": str(root / "test.db"),
                "UPLOAD_FOLDER": str(root / "uploads"),
                "EXPORT_FOLDER": str(root / "exports"),
            }
        )
        self.context = self.app.app_context()
        self.context.push()

    def tearDown(self):
        self.context.pop()
        self.tempdir.cleanup()

    def _prepare_paired_subject_line_case(self, connection, track: str):
        seed_demo_data(connection)
        exam = connection.execute(
            "SELECT id, exam_key FROM exams ORDER BY exam_date DESC LIMIT 1"
        ).fetchone()
        result_ids = [
            row["id"]
            for row in connection.execute(
                """
                SELECT er.id
                FROM exam_results er
                JOIN student_exam_enrollments se ON se.id=er.enrollment_id
                WHERE er.exam_id=? AND se.track=?
                ORDER BY er.id
                """,
                (exam["id"], track),
            )
        ]
        for index, result_id in enumerate(result_ids):
            total = float(700 - (index % 5) * 20)
            connection.execute(
                "UPDATE exam_results SET raw_total=?, scaled_total=?, effective_total=? WHERE id=?",
                (total, total, total, result_id),
            )
            scores = (
                {"chemistry": 72 + index % 4, "politics": 42 + index % 4,
                 "biology": 82 + index % 4, "geography": 32 + index % 4}
                if track == "物理"
                else {"politics": 76 + index % 4, "chemistry": 46 + index % 4,
                      "geography": 86 + index % 4, "biology": 36 + index % 4}
            )
            for code, score in scores.items():
                connection.execute(
                    """
                    INSERT INTO subject_scores(
                      exam_result_id, subject_code, raw_score, scaled_score, effective_score
                    ) VALUES (?, ?, ?, ?, ?)
                    ON CONFLICT(exam_result_id, subject_code) DO UPDATE SET
                      raw_score=excluded.raw_score,
                      scaled_score=excluded.scaled_score,
                      effective_score=excluded.effective_score
                    """,
                    (result_id, code, score, score, score),
                )
        for line_type, value in (("特控", 500), ("211", 600), ("985", 650)):
            connection.execute(
                """
                UPDATE score_lines SET line_value=?
                WHERE exam_id=? AND track=? AND line_type=? AND subject_code='__TOTAL__'
                """,
                (value, exam["id"], track, line_type),
            )
        connection.commit()
        return exam, result_ids

    def _expected_subject_line(self, connection, result_ids, code: str, total_line: float):
        placeholders = ",".join("?" for _ in result_ids)
        totals = [
            Decimal(str(row["scaled_total"]))
            for row in connection.execute(
                f"SELECT scaled_total FROM exam_results WHERE id IN ({placeholders})",
                result_ids,
            )
            if row["scaled_total"] is not None
            and Decimal(str(row["scaled_total"])) >= Decimal(str(total_line))
        ]
        score_field = (
            "scaled_score"
            if code in {"chemistry", "biology", "politics", "geography"}
            else "effective_score"
        )
        scores = [
            Decimal(str(row[score_field]))
            for row in connection.execute(
                f"""
                SELECT ss.scaled_score, ss.effective_score, er.scaled_total
                FROM subject_scores ss
                JOIN exam_results er ON er.id=ss.exam_result_id
                WHERE ss.exam_result_id IN ({placeholders})
                  AND ss.subject_code=? AND ss.{score_field} IS NOT NULL
                """,
                [*result_ids, code],
            )
            if row["scaled_total"] is not None
            and Decimal(str(row["scaled_total"])) >= Decimal(str(total_line))
        ]
        raw = (
            (sum(scores) / Decimal(len(scores)))
            * Decimal(str(total_line))
            / (sum(totals) / Decimal(len(totals)))
        )
        return _mround_half(raw)

    def test_new_base_information_template_aliases_are_saved_and_reusable(self):
        path = Path(self.tempdir.name) / "base.xlsx"
        workbook = Workbook()
        school = workbook.active
        school.title = "学校信息"
        school.append(["省份", "学校名称", "年级", "班级", "班级人数", "班级特控目标人数"])
        school.append(["重庆", "测试中学", "高2027届", "1", "30", "18"])
        students = workbook.create_sheet("学生信息")
        students.append(["学生姓名", "学校", "年级", "班级", "学号", "选科标签"])
        students.append(["测试学生", "测试中学", "高2027届", "1", "S001", ""])
        workbook.save(path)
        payload = parse_base_information_workbook(path)
        self.assertEqual([issue for issue in payload.issues if issue.severity == "error"], [])
        self.assertEqual(payload.schools[0]["control_target"], 18)
        self.assertEqual(payload.students[0]["subject_tag"], "")
        result = save_base_information(get_db(), payload)
        self.assertEqual(result["students"], 1)
        reference = Path(self.tempdir.name) / "reusable.xlsx"
        build_reference_workbook_from_base_information(get_db(), reference)
        rebuilt = load_workbook(reference, data_only=True)
        self.assertEqual(rebuilt.sheetnames, ["学校信息", "学生信息"])
        self.assertEqual(rebuilt["学生信息"][2][3].value, "测试学生")

    def test_dynamic_stage_targets_keep_order_persist_and_restore(self):
        path = Path(self.tempdir.name) / "stage_targets.xlsx"
        workbook = Workbook()
        school = workbook.active
        school.title = "学校信息"
        school.append(["省份", "学校", "年级", "班级", "方向", "班级人数", "一诊目标", "二诊目标", "高考目标"])
        school.append(["重庆", "阶段目标中学", "高2028届", "1", "物理", 30, 12, 15, 18])
        students = workbook.create_sheet("学生信息")
        students.append(["学校", "年级", "班级", "学生姓名", "学号", "选科标签"])
        students.append(["阶段目标中学", "高2028届", "1", "阶段学生", "T001", "物化生"])
        workbook.save(path)
        payload = parse_base_information_workbook(path)
        self.assertEqual(payload.schools[0]["targets"], [
            {"name": "一诊", "order": 0, "count": 12},
            {"name": "二诊", "order": 1, "count": 15},
            {"name": "高考", "order": 2, "count": 18},
        ])
        save_base_information(get_db(), payload)
        saved = get_db().execute(
            """SELECT target_name, target_order, target_count FROM base_class_targets
               ORDER BY target_order"""
        ).fetchall()
        self.assertEqual([tuple(row) for row in saved], [("一诊", 0, 12), ("二诊", 1, 15), ("高考", 2, 18)])
        formal = get_db().execute(
            "SELECT target_name, target_order, target_count FROM class_stage_targets ORDER BY target_order"
        ).fetchall()
        self.assertEqual([tuple(row) for row in formal], [("一诊", 0, 12), ("二诊", 1, 15), ("高考", 2, 18)])
        grade_id = get_db().execute("SELECT id FROM grades WHERE name='高2028届'").fetchone()["id"]
        get_db().execute(
            "INSERT INTO exams(exam_key, grade_id, name, exam_date) VALUES ('stage-exam', ?, '高三一诊考试', '2027-09-01')",
            (grade_id,),
        )
        get_db().commit()
        filters = filter_options(get_db(), "stage-exam")
        self.assertEqual(filters["targets"], ["一诊", "二诊", "高考"])
        self.assertEqual(filters["default_target_name"], "高考")
        get_db().execute("UPDATE exams SET name='普通月考' WHERE exam_key='stage-exam'")
        get_db().commit()
        self.assertEqual(filter_options(get_db(), "stage-exam")["default_target_name"], "高考")
        reference = Path(self.tempdir.name) / "stage_targets_rebuilt.xlsx"
        build_reference_workbook_from_base_information(get_db(), reference)
        rebuilt = load_workbook(reference, data_only=True)
        self.assertEqual([cell.value for cell in rebuilt["学校信息"][1][6:9]], ["一诊目标", "二诊目标", "高考目标"])
        rebuilt.close()

    def test_invalid_non_integer_stage_target_is_blocked(self):
        path = Path(self.tempdir.name) / "invalid_stage_target.xlsx"
        workbook = Workbook()
        school = workbook.active
        school.title = "学校信息"
        school.append(["省份", "学校", "年级", "班级", "方向", "班级人数", "一诊目标"])
        school.append(["重庆", "测试中学", "高2028届", "1", "物理", 30, 12.5])
        students = workbook.create_sheet("学生信息")
        students.append(["学校", "年级", "班级", "学生姓名", "学号", "选科标签"])
        students.append(["测试中学", "高2028届", "1", "测试学生", "T001", "物化生"])
        workbook.save(path)
        payload = parse_base_information_workbook(path)
        self.assertTrue(any(issue.issue_code == "target_count_invalid" for issue in payload.issues))

    def test_base_target_snapshot_does_not_change_exam_scores(self):
        connection = get_db()
        seed_demo_data(connection)
        before_results = tuple(connection.execute(
            """SELECT COUNT(*), SUM(COALESCE(raw_total, 0)),
                      SUM(COALESCE(scaled_total, 0)), SUM(COALESCE(effective_total, 0))
               FROM exam_results"""
        ).fetchone())
        before_subjects = tuple(connection.execute(
            """SELECT COUNT(*), SUM(COALESCE(raw_score, 0)),
                      SUM(COALESCE(scaled_score, 0)), SUM(COALESCE(effective_score, 0))
               FROM subject_scores"""
        ).fetchone())
        reference = connection.execute(
            """SELECT o.name AS school, g.name AS grade, c.class_name, c.track,
                      s.name AS student_name,
                      COALESCE(s.display_student_no, s.student_no) AS student_no
               FROM student_exam_enrollments se
               JOIN organizations o ON o.id=se.organization_id
               JOIN grades g ON g.id=se.grade_id
               JOIN classes c ON c.id=se.class_id
               JOIN students s ON s.id=se.student_id
               LIMIT 1"""
        ).fetchone()
        path = Path(self.tempdir.name) / "target_only_base_update.xlsx"
        workbook = Workbook()
        school = workbook.active
        school.title = "学校信息"
        school.append(["省份", "学校", "年级", "班级", "方向", "班级人数", "特控目标", "一诊目标"])
        school.append(["", reference["school"], reference["grade"], reference["class_name"], reference["track"], 50, 42, 40])
        students = workbook.create_sheet("学生信息")
        students.append(["学校", "年级", "班级", "学生姓名", "学号", "选科标签"])
        students.append([
            reference["school"], reference["grade"], reference["class_name"],
            reference["student_name"], reference["student_no"], "",
        ])
        workbook.save(path)
        save_base_information(connection, parse_base_information_workbook(path))
        after_results = tuple(connection.execute(
            """SELECT COUNT(*), SUM(COALESCE(raw_total, 0)),
                      SUM(COALESCE(scaled_total, 0)), SUM(COALESCE(effective_total, 0))
               FROM exam_results"""
        ).fetchone())
        after_subjects = tuple(connection.execute(
            """SELECT COUNT(*), SUM(COALESCE(raw_score, 0)),
                      SUM(COALESCE(scaled_score, 0)), SUM(COALESCE(effective_score, 0))
               FROM subject_scores"""
        ).fetchone())
        self.assertEqual(after_results, before_results)
        self.assertEqual(after_subjects, before_subjects)


    def test_preparation_reuses_saved_base_without_temporary_reference_workbook(self):
        source = Path(self.tempdir.name) / "source.xlsx"
        output = Path(self.tempdir.name) / "prepared.xlsx"
        self._make_valid_workbook(source)
        save_base_information(get_db(), parse_base_information_workbook(source))
        result = prepare_import_workbook(
            None, source, output, "高2028届", "高二下期末考试", "2027年06月",
            reference_data=saved_base_reference_data(get_db()),
        )
        self.assertEqual(result["matched_rows"], 1)
        self.assertIn("基础信息写入", result["timings"])
        payload = parse_workbook(output, get_db())
        self.assertEqual([issue for issue in payload.issues if issue.severity == "error"], [])

    def test_base_information_is_a_grade_snapshot_and_preserves_class_track(self):
        path = Path(self.tempdir.name) / "base_snapshot.xlsx"
        self._make_valid_workbook(path)
        save_base_information(get_db(), parse_base_information_workbook(path))
        get_db().execute(
            """INSERT INTO base_students(
                 school_name, grade_name, class_name, student_name, student_no
               ) VALUES ('测试中学', '高2027届', '1', '测试学生', 'STALE')"""
        )
        get_db().commit()
        workbook = load_workbook(path)
        workbook["学校基础信息表"].append(["重庆", "测试中学", "高2027届", "9", "历史", 30, 5])
        workbook.save(path)
        save_base_information(get_db(), parse_base_information_workbook(path))
        self.assertIsNone(get_db().execute("SELECT id FROM base_students WHERE student_no='STALE'").fetchone())
        self.assertIsNotNone(
            get_db().execute(
                """SELECT c.id FROM classes c JOIN organizations o ON o.id=c.organization_id
                   JOIN grades g ON g.id=c.grade_id
                   WHERE o.name='测试中学' AND g.name='高2027届'
                     AND c.class_name='9' AND c.track='历史' AND c.archived_at IS NULL"""
            ).fetchone()
        )

    def test_three_separate_grade_base_uploads_retain_earlier_grades(self):
        connection = get_db()
        grades = ["高2027届", "高2028届", "高2029届"]
        for index, grade in enumerate(grades, start=1):
            path = Path(self.tempdir.name) / f"base_{index}.xlsx"
            self._make_base_grades_workbook(path, [grade])
            save_base_information(connection, parse_base_information_workbook(path))
        self.assertEqual(
            [row["grade_name"] for row in connection.execute(
                "SELECT DISTINCT grade_name FROM base_students ORDER BY grade_name"
            )],
            grades,
        )
        self.assertEqual(
            connection.execute("SELECT COUNT(*) AS n FROM base_school_classes").fetchone()["n"],
            3,
        )

    def test_combined_three_grade_base_upload_keeps_same_class_numbers_separate(self):
        connection = get_db()
        grades = ["高2027届", "高2028届", "高2029届"]
        path = Path(self.tempdir.name) / "base_all_grades.xlsx"
        self._make_base_grades_workbook(path, grades)
        save_base_information(connection, parse_base_information_workbook(path))
        classes = connection.execute(
            """SELECT c.id, g.name AS grade_name
               FROM classes c JOIN grades g ON g.id=c.grade_id
               WHERE c.class_name='1' AND c.archived_at IS NULL
               ORDER BY g.name"""
        ).fetchall()
        self.assertEqual([row["grade_name"] for row in classes], grades)
        self.assertEqual(len({row["id"] for row in classes}), 3)
        self.assertEqual(
            connection.execute("SELECT COUNT(*) AS n FROM base_students").fetchone()["n"],
            3,
        )

    def test_demo_analytics_invariants(self):
        connection = get_db()
        seed_demo_data(connection)
        students = connection.execute("SELECT student_no, name FROM students WHERE student_no LIKE 'DEMO%'").fetchall()
        self.assertEqual(len(students), 120)
        for student in students:
            self.assertEqual(student['name'], f"演示学生{int(student['student_no'][4:]):03d}")
        exam = connection.execute("SELECT exam_key FROM exams ORDER BY exam_date DESC LIMIT 1").fetchone()
        scope = Scope(exam["exam_key"], "物理")
        overall = overview(connection, scope, 40)
        self.assertGreater(overall["reference_count"], 0)
        self.assertEqual(sum(item["count"] for item in overall["student_types"]), overall["reference_count"])
        self.assertEqual([item["name"] for item in overall["student_types"]], ["学生"])
        self.assertGreaterEqual(len(overall["history"]), 1)
        options = filter_options(connection, exam["exam_key"])
        self.assertNotIn("不支持的类别", [item["name"] for item in options["organizations"]])
        self.assertNotIn("不支持的类别", [item["class_name"] for item in options["classes"]])
        self.assertLessEqual(overall["lines"]["特控"]["count"], overall["reference_count"])
        self.assertEqual(overall["standard_deviation_type"], "population")
        bands = score_bands(connection, scope)
        self.assertTrue(bands["reconcile"]["matches"])
        self.assertEqual(sum(item["count"] for item in bands["line_bands"]), overall["reference_count"])
        subjects = subject_metrics(connection, scope)
        self.assertGreaterEqual(len(subjects["subjects"]), 6)
        for item in subjects["subjects"]:
            if item["both_online_count"] is not None:
                self.assertLessEqual(item["both_online_count"], item["control_online_count"])
        groups = student_groups(connection, scope, 40)
        self.assertGreaterEqual(groups["common_student_count"], len(groups["converted"]))

    def test_current_critical_students_include_previous_total_when_available(self):
        connection = get_db()
        seed_demo_data(connection)
        exams = connection.execute(
            "SELECT id, exam_key FROM exams ORDER BY exam_date, id"
        ).fetchall()
        scope = Scope(exams[-1]["exam_key"], "物理")
        exam_keys = [exam["exam_key"] for exam in exams]

        groups = student_groups(connection, scope, 40, exam_keys)
        self.assertTrue(groups["current_critical"])
        previous_exam_id = groups["previous_exam"]["id"]
        previous_scores = {
            row["identity_key"]: row["effective_total"]
            for row in connection.execute(
                """
                SELECT COALESCE(s.identity_key, 'legacy|' || s.student_no) AS identity_key,
                       er.effective_total
                FROM exam_results er
                JOIN students s ON s.id=er.student_id
                WHERE er.exam_id=?
                  AND er.attendance_status='normal' AND er.effective_total IS NOT NULL
                """,
                (previous_exam_id,),
            )
        }
        comparable = [
            item for item in groups["current_critical"]
            if item["identity_key"] in previous_scores
        ]
        self.assertTrue(comparable)
        for item in comparable:
            self.assertEqual(item["previous_total"], previous_scores[item["identity_key"]])

        missing_item = comparable[0]
        student_id = connection.execute(
            "SELECT id FROM students WHERE COALESCE(identity_key, 'legacy|' || student_no)=?",
            (missing_item["identity_key"],),
        ).fetchone()["id"]
        connection.execute(
            "DELETE FROM exam_results WHERE exam_id=? AND student_id=?",
            (previous_exam_id, student_id),
        )
        connection.commit()

        refreshed = student_groups(connection, scope, 40, exam_keys)
        refreshed_item = next(
            item for item in refreshed["current_critical"]
            if item["identity_key"] == missing_item["identity_key"]
        )
        self.assertNotIn("previous_total", refreshed_item)


    def test_import_preview_and_commit(self):
        workbook_path = Path(self.tempdir.name) / "valid_import.xlsx"
        self._make_valid_workbook(workbook_path)
        connection = get_db()
        preview = preview_workbook(connection, workbook_path, workbook_path.name)
        self.assertEqual(preview["error_count"], 0, preview["issues"])
        result = commit_batch(connection, preview["batch_key"], "reject")
        self.assertEqual(result["committed_rows"], 1)
        self.assertEqual(len(result["exam_keys"]), 1)
        stored = connection.execute("SELECT effective_total, total_source FROM exam_results").fetchone()
        self.assertEqual(stored["effective_total"], 540)
        self.assertEqual(stored["total_source"], "scaled")
        exam_key = connection.execute("SELECT exam_key FROM exams").fetchone()["exam_key"]
        imported_overview = overview(connection, Scope(exam_key, "物理"), target_name="特控")
        self.assertEqual(imported_overview["targets"]["configured_classes"], 1)
        self.assertEqual(imported_overview["targets"]["met"], 1)

    def test_partial_import_keeps_valid_rows_and_exposes_all_issue_details(self):
        workbook_path = Path(self.tempdir.name) / "partial_import.xlsx"
        self._make_valid_workbook(workbook_path)
        workbook = load_workbook(workbook_path)
        workbook["学生信息"].append(
            ["问题学生", "S002", "高2027届", "1", "测试中学", "物化生"]
        )
        physics = workbook["物理方向学生成绩"]
        invalid_row = list(
            physics.iter_rows(min_row=2, max_row=2, max_col=48, values_only=True)
        )[0]
        invalid_row = list(invalid_row)
        invalid_row[3] = "S002"
        invalid_row[5] = "问题学生"
        invalid_row[45] = 900
        physics.append(invalid_row)
        workbook.save(workbook_path)

        connection = get_db()
        preview = preview_workbook(connection, workbook_path, workbook_path.name)
        self.assertGreater(preview["error_count"], 0)
        self.assertGreater(preview["warning_count"], 0)
        self.assertEqual(preview["summary"]["score_rows"], 2)
        self.assertEqual(preview["valid_row_count"], 1)
        self.assertEqual(preview["blocked_row_count"], 1)

        stored_preview = get_batch_preview(connection, preview["batch_key"])
        self.assertEqual(stored_preview["valid_row_count"], 1)
        errors = list_batch_issues(connection, preview["batch_key"], "error", 1, 1)
        self.assertGreaterEqual(errors["total"], 1)
        self.assertEqual(len(errors["items"]), 1)
        self.assertIn("message", errors["items"][0])
        self.assertIn("suggestion", errors["items"][0])
        client = self.app.test_client()
        self.assertEqual(
            client.get(f"/api/import/{preview['batch_key']}").status_code, 200
        )
        issue_response = client.get(
            f"/api/import/{preview['batch_key']}/issues?severity=error&page=1&page_size=1"
        )
        self.assertEqual(issue_response.status_code, 200)
        self.assertEqual(len(issue_response.get_json()["items"]), 1)
        unconfirmed_response = client.post(
            f"/api/import/{preview['batch_key']}/commit",
            json={"mode": "reject", "allow_partial": True, "warnings_confirmed": False},
        )
        self.assertEqual(unconfirmed_response.status_code, 400)
        self.assertIn("确认", unconfirmed_response.get_json()["error"])

        with self.assertRaisesRegex(ValueError, "阻止入库"):
            commit_batch(connection, preview["batch_key"], "reject")
        with self.assertRaisesRegex(ValueError, "确认"):
            commit_batch(
                connection,
                preview["batch_key"],
                "reject",
                allow_partial=True,
                warnings_confirmed=False,
            )

        result = commit_batch(
            connection,
            preview["batch_key"],
            "reject",
            allow_partial=True,
            warnings_confirmed=True,
        )
        self.assertEqual(result["status"], "committed_partial")
        self.assertEqual(result["committed_rows"], 1)
        self.assertEqual(result["blocked_rows_skipped"], 1)
        self.assertEqual(
            connection.execute("SELECT COUNT(*) AS n FROM exam_results").fetchone()["n"],
            1,
        )
        self.assertEqual(
            connection.execute("SELECT name FROM students").fetchone()["name"],
            "测试学生",
        )

    def test_two_file_preparation_matches_identity_and_locks_exam_context(self):
        source = Path(self.tempdir.name) / "source.xlsx"
        output = Path(self.tempdir.name) / "prepared.xlsx"
        self._make_valid_workbook(source)
        workbook = load_workbook(source)
        extra_sheet = workbook.create_sheet("其他学生数据", 4)
        extra_sheet.append(["年级", "考试名称", "考试时间", "学生类型", "班级", "姓名", "物理", "历史", "总分"])
        workbook.save(source)

        prepared = prepare_import_workbook(
            source, source, output, "高2028届", "高二下期末考试", "2027年06月"
        )

        self.assertEqual(prepared["matched_rows"], 1)
        self.assertEqual(prepared["unmatched_rows"], 0)
        workbook = load_workbook(output, data_only=True)
        physics = workbook["物理方向学生成绩"]
        self.assertEqual(
            [physics.cell(1, column).value for column in range(1, 5)],
            ["班级", "姓名", "语文", "语文排名"],
        )
        self.assertEqual(
            [physics.cell(2, column).value for column in (1, 2)],
            ["1", "测试学生"],
        )
        self.assertEqual(
            [physics.cell(1, column).value for column in range(35, 40)],
            ["年级", "考试名称", "考试时间", "学校", "学号"],
        )
        self.assertTrue(physics.column_dimensions["AI"].hidden)
        workbook.close()
        payload = parse_workbook(output, get_db())
        self.assertEqual([issue for issue in payload.issues if issue.severity == "error"], [])
        self.assertEqual({score["exam_key"] for score in payload.scores}, {
            payload.scores[0]["exam_key"]
        })

    def test_two_file_preparation_accepts_missing_optional_score_line_sheet(self):
        source = Path(self.tempdir.name) / "source_without_lines.xlsx"
        output = Path(self.tempdir.name) / "prepared_without_lines.xlsx"
        self._make_valid_workbook(source)
        workbook = load_workbook(source)
        workbook.create_sheet("其他学生数据").append(
            ["年级", "考试名称", "考试时间", "学生类型", "班级", "姓名", "物理", "历史", "总分"]
        )
        del workbook["分数线表"]
        workbook.save(source)

        prepared = prepare_import_workbook(
            source, source, output, "高2028届", "高二下期末考试", "2027年06月"
        )
        self.assertEqual(prepared["matched_rows"], 1)
        workbook = load_workbook(output, read_only=True)
        self.assertNotIn("分数线表", workbook.sheetnames)
        workbook.close()

    def test_manual_subject_lines_override_derived_lines(self):
        connection = get_db()
        seed_demo_data(connection)
        exam_key = connection.execute(
            "SELECT exam_key FROM exams ORDER BY exam_date DESC LIMIT 1"
        ).fetchone()["exam_key"]
        client = self.app.test_client()
        response = client.put(
            "/api/settings/score-lines",
            json={
                "exam_key": exam_key,
                "track": "物理",
                "特控": 500,
                "211": 600,
                "985": 650,
                "subject_lines": {
                    "chinese": {"特控": 91, "211": 101, "985": 111}
                },
            },
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        subjects = subject_metrics(connection, Scope(exam_key, "物理"), include_history=False)
        chinese = next(item for item in subjects["subjects"] if item["code"] == "chinese")
        self.assertEqual(chinese["lines"], {"特控": 91.0, "211": 101.0, "985": 111.0})

    def test_subject_online_rates_use_all_students_in_scope_as_denominator(self):
        connection = get_db()
        seed_demo_data(connection)
        exam = connection.execute(
            "SELECT id, exam_key FROM exams ORDER BY exam_date DESC LIMIT 1"
        ).fetchone()
        result_ids = [
            row["id"]
            for row in connection.execute(
                """
                SELECT er.id
                FROM exam_results er
                JOIN student_exam_enrollments se ON se.id=er.enrollment_id
                WHERE er.exam_id=? AND se.track='物理' AND er.attendance_status='normal'
                  AND er.effective_total IS NOT NULL
                ORDER BY er.id
                """,
                (exam["id"],),
            )
        ]
        self.assertGreater(len(result_ids), 1)
        placeholders = ",".join("?" for _ in result_ids)
        connection.execute(
            f"DELETE FROM subject_scores WHERE subject_code='chinese' AND exam_result_id IN ({placeholders})",
            result_ids,
        )
        connection.execute(
            """
            INSERT INTO subject_scores(
              exam_result_id, subject_code, raw_score, scaled_score, effective_score
            ) VALUES (?, 'chinese', 100, NULL, 100)
            """,
            (result_ids[0],),
        )
        connection.execute(
            """
            INSERT INTO score_lines(
              exam_id, track, line_type, subject_code, line_value, source, rule_version
            ) VALUES (?, '物理', '特控', 'chinese', 90, 'manual', 'test')
            ON CONFLICT(exam_id, track, line_type, subject_code) DO UPDATE SET line_value=excluded.line_value
            """,
            (exam["id"],),
        )
        connection.commit()

        subjects = subject_metrics(
            connection, Scope(exam["exam_key"], "物理"), include_history=False
        )
        chinese = next(item for item in subjects["subjects"] if item["code"] == "chinese")
        self.assertEqual(chinese["reference_count"], 1)
        self.assertEqual(chinese["control_online_count"], 1)
        self.assertAlmostEqual(chinese["control_online_rate"], 1 / len(result_ids))
        self.assertAlmostEqual(chinese["online"]["特控"]["rate"], 1 / len(result_ids))
        class_id = connection.execute(
            "SELECT class_id FROM student_exam_enrollments se JOIN exam_results er ON er.enrollment_id=se.id WHERE er.id=?",
            (result_ids[0],),
        ).fetchone()["class_id"]
        classes = class_metrics(
            connection,
            Scope(exam["exam_key"], "物理", class_id=class_id),
            include_history=False,
        )
        class_chinese = next(
            item for item in classes["subject_comparison"] if item["code"] == "chinese"
        )
        class_reference_count = classes["classes"][0]["reference_count"]
        self.assertEqual(class_chinese["reference_count"], 1)
        self.assertEqual(class_chinese["control_count"], 1)
        self.assertAlmostEqual(class_chinese["control_rate"], 1 / class_reference_count)

    def test_score_lines_save_incrementally_and_survive_app_restart(self):
        connection = get_db()
        seed_demo_data(connection)
        exam = connection.execute(
            "SELECT id, exam_key FROM exams ORDER BY exam_date DESC LIMIT 1"
        ).fetchone()
        connection.execute(
            "DELETE FROM score_lines WHERE exam_id=? AND track='历史' AND subject_code='__TOTAL__'",
            (exam["id"],),
        )
        connection.commit()
        client = self.app.test_client()

        expected = {}
        for line_type, value in (("特控", 500), ("211", 570), ("985", 630)):
            response = client.put(
                "/api/settings/score-lines",
                json={"exam_key": exam["exam_key"], "track": "历史", line_type: value},
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            expected[line_type] = value
            saved = _line_map(connection, exam["id"], "历史")
            for saved_type, saved_value in expected.items():
                self.assertEqual(saved[saved_type], saved_value)

        version_before_blank = connection.execute(
            "SELECT data_version FROM exams WHERE id=?", (exam["id"],)
        ).fetchone()["data_version"]
        blank = client.put(
            "/api/settings/score-lines",
            json={
                "exam_key": exam["exam_key"], "track": "历史",
                "本科": "", "特控": "", "211": None, "985": "   ",
            },
        )
        self.assertEqual(blank.status_code, 200, blank.get_json())
        self.assertEqual(blank.get_json()["status"], "unchanged")
        self.assertEqual(
            connection.execute("SELECT data_version FROM exams WHERE id=?", (exam["id"],)).fetchone()["data_version"],
            version_before_blank,
        )

        benchmark = client.get(f"/api/analytics/benchmark?exam_key={exam['exam_key']}")
        self.assertEqual(benchmark.status_code, 200, benchmark.get_json())
        self.assertEqual(
            benchmark.get_json()["tracks"]["历史"]["total_lines"],
            {"本科": None, "特控": 500.0, "211": 570.0, "985": 630.0},
        )

        second_app = create_app(
            {
                "TESTING": True,
                "DATABASE": self.app.config["DATABASE"],
                "UPLOAD_FOLDER": self.app.config["UPLOAD_FOLDER"],
                "EXPORT_FOLDER": self.app.config["EXPORT_FOLDER"],
            }
        )
        with second_app.app_context():
            reloaded = second_app.test_client().get(
                f"/api/analytics/benchmark?exam_key={exam['exam_key']}"
            )
            self.assertEqual(reloaded.status_code, 200, reloaded.get_json())
            self.assertEqual(
                reloaded.get_json()["tracks"]["历史"]["total_lines"],
                {"本科": None, "特控": 500.0, "211": 570.0, "985": 630.0},
            )

    def test_undergraduate_and_subject_lines_allow_partial_updates(self):
        connection = get_db()
        seed_demo_data(connection)
        exam = connection.execute(
            "SELECT id, exam_key FROM exams ORDER BY exam_date DESC LIMIT 1"
        ).fetchone()
        connection.execute(
            "DELETE FROM score_lines WHERE exam_id=? AND track='物理'",
            (exam["id"],),
        )
        connection.commit()
        client = self.app.test_client()

        undergraduate = client.put(
            "/api/settings/score-lines",
            json={"exam_key": exam["exam_key"], "track": "物理", "本科": 450},
        )
        self.assertEqual(undergraduate.status_code, 200, undergraduate.get_json())
        control = client.put(
            "/api/settings/score-lines",
            json={
                "exam_key": exam["exam_key"], "track": "物理", "特控": 500,
                "subject_lines": {"chinese": {"特控": 90, "211": "", "985": None}},
            },
        )
        self.assertEqual(control.status_code, 200, control.get_json())
        subject_211 = client.put(
            "/api/settings/score-lines",
            json={
                "exam_key": exam["exam_key"], "track": "物理",
                "subject_lines": {"chinese": {"特控": "", "211": 100}},
            },
        )
        self.assertEqual(subject_211.status_code, 200, subject_211.get_json())
        saved_subject_lines = {
            row["line_type"]: row["line_value"]
            for row in connection.execute(
                """
                SELECT line_type, line_value FROM score_lines
                WHERE exam_id=? AND track='物理' AND subject_code='chinese'
                """,
                (exam["id"],),
            )
        }
        self.assertEqual(saved_subject_lines, {"特控": 90.0, "211": 100.0})
        self.assertEqual(
            _line_map(connection, exam["id"], "物理"),
            {"本科": 450.0, "特控": 500.0, "211": None, "985": None},
        )

    def test_single_total_line_only_derives_matching_subject_line(self):
        connection = get_db()
        seed_demo_data(connection)
        exam = connection.execute(
            "SELECT id, exam_key FROM exams ORDER BY exam_date DESC LIMIT 1"
        ).fetchone()
        connection.execute(
            "DELETE FROM score_lines WHERE exam_id=? AND track='物理'",
            (exam["id"],),
        )
        connection.commit()
        response = self.app.test_client().put(
            "/api/settings/score-lines",
            json={"exam_key": exam["exam_key"], "track": "物理", "特控": 500},
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        lines = _derived_subject_lines(connection, exam["exam_key"], "物理")
        self.assertIsNotNone(lines["chinese"]["特控"])
        self.assertIsNone(lines["chinese"]["211"])
        self.assertIsNone(lines["chinese"]["985"])
        self.assertEqual(lines["chemistry"], lines["politics"])
        self.assertEqual(lines["biology"], lines["geography"])

    def test_excel_mround_half_matches_half_up_boundaries(self):
        self.assertEqual(_mround_half(89.6729635090305), 89.5)
        self.assertEqual(_mround_half(89.75), 90.0)
        self.assertEqual(_mround_half(89.25), 89.5)

    def test_subject_line_keeps_full_precision_until_final_mround(self):
        connection = get_db()
        exam, result_ids = self._prepare_paired_subject_line_case(connection, "物理")
        self.assertGreaterEqual(len(result_ids), 2)
        placeholders = ",".join("?" for _ in result_ids)
        connection.execute(
            f"UPDATE exam_results SET scaled_total=400 WHERE id IN ({placeholders})",
            result_ids,
        )
        connection.execute(
            "UPDATE exam_results SET scaled_total=500 WHERE id IN (?, ?)",
            result_ids[:2],
        )
        for result_id, score in zip(result_ids[:2], (89.748, 89.75)):
            connection.execute(
                """
                INSERT INTO subject_scores(
                  exam_result_id, subject_code, raw_score, scaled_score, effective_score
                ) VALUES (?, 'chinese', ?, NULL, ?)
                ON CONFLICT(exam_result_id, subject_code) DO UPDATE SET
                  raw_score=excluded.raw_score,
                  scaled_score=excluded.scaled_score,
                  effective_score=excluded.effective_score
                """,
                (result_id, score, score),
            )
        connection.execute(
            "DELETE FROM score_lines WHERE exam_id=? AND track='物理'",
            (exam["id"],),
        )
        connection.execute(
            """
            INSERT INTO score_lines(
              exam_id, track, line_type, subject_code, line_value, source, rule_version
            ) VALUES (?, '物理', '特控', '__TOTAL__', 500, 'manual', 'test')
            """,
            (exam["id"],),
        )
        connection.commit()

        lines = _derived_subject_lines(connection, exam["exam_key"], "物理")
        self.assertEqual(lines["chinese"]["特控"], 89.5)
        self.assertIsNone(lines["chinese"]["211"])
        self.assertIsNone(lines["chinese"]["985"])

    def test_subject_lines_use_scaled_total_and_scaled_elective_scores(self):
        connection = get_db()
        exam, result_ids = self._prepare_paired_subject_line_case(connection, "物理")
        self.assertGreaterEqual(len(result_ids), 2)
        placeholders = ",".join("?" for _ in result_ids)
        connection.execute(
            f"""
            UPDATE exam_results
            SET scaled_total=400, effective_total=700
            WHERE id IN ({placeholders})
            """,
            result_ids,
        )
        connection.execute(
            "UPDATE exam_results SET scaled_total=600, effective_total=400 WHERE id=?",
            (result_ids[0],),
        )
        for result_id, chinese, chemistry_scaled, chemistry_effective in (
            (result_ids[0], 90, 80, 10),
            (result_ids[1], 10, 20, 99),
        ):
            connection.execute(
                """
                INSERT INTO subject_scores(
                  exam_result_id, subject_code, raw_score, scaled_score, effective_score
                ) VALUES (?, 'chinese', ?, NULL, ?)
                ON CONFLICT(exam_result_id, subject_code) DO UPDATE SET
                  raw_score=excluded.raw_score,
                  scaled_score=excluded.scaled_score,
                  effective_score=excluded.effective_score
                """,
                (result_id, chinese, chinese),
            )
            connection.execute(
                """
                INSERT INTO subject_scores(
                  exam_result_id, subject_code, raw_score, scaled_score, effective_score
                ) VALUES (?, 'chemistry', ?, ?, ?)
                ON CONFLICT(exam_result_id, subject_code) DO UPDATE SET
                  raw_score=excluded.raw_score,
                  scaled_score=excluded.scaled_score,
                  effective_score=excluded.effective_score
                """,
                (result_id, chemistry_effective, chemistry_scaled, chemistry_effective),
            )
        connection.execute(
            "DELETE FROM score_lines WHERE exam_id=? AND track='物理'",
            (exam["id"],),
        )
        connection.execute(
            """
            INSERT INTO score_lines(
              exam_id, track, line_type, subject_code, line_value, source, rule_version
            ) VALUES (?, '物理', '特控', '__TOTAL__', 500, 'manual', 'test')
            """,
            (exam["id"],),
        )
        connection.commit()

        lines = _derived_subject_lines(connection, exam["exam_key"], "物理")
        self.assertEqual(lines["chinese"]["特控"], 75.0)
        self.assertEqual(lines["chemistry"]["特控"], 66.5)
        self.assertEqual(lines["politics"]["特控"], 66.5)

    def test_partial_line_order_is_checked_against_saved_values(self):
        connection = get_db()
        seed_demo_data(connection)
        exam = connection.execute(
            "SELECT id, exam_key FROM exams ORDER BY exam_date DESC LIMIT 1"
        ).fetchone()
        connection.execute(
            "DELETE FROM score_lines WHERE exam_id=? AND track='历史'",
            (exam["id"],),
        )
        connection.commit()
        client = self.app.test_client()
        self.assertEqual(
            client.put(
                "/api/settings/score-lines",
                json={"exam_key": exam["exam_key"], "track": "历史", "特控": 500},
            ).status_code,
            200,
        )
        invalid_total = client.put(
            "/api/settings/score-lines",
            json={"exam_key": exam["exam_key"], "track": "历史", "211": 490},
        )
        self.assertEqual(invalid_total.status_code, 400)
        self.assertIsNone(_line_map(connection, exam["id"], "历史")["211"])

    def test_frontend_persists_only_exam_and_track_selection_for_lines(self):
        source = (Path(self.app.root_path) / "static" / "app.js").read_text(encoding="utf-8")
        self.assertIn('"student-analysis-selected-exam"', source)
        self.assertIn('"student-analysis-selected-track"', source)
        self.assertIn("persistSelectedScope();", source)
        self.assertNotIn('value="${trackData.total_lines[lineType] ?? ""}" required', source)
        self.assertNotRegex(
            source,
            r"localStorage\.(?:setItem|getItem)\([^\n]*(?:total_lines|line_value|特控|211|985|本科)",
        )

    def test_successful_import_opens_benchmark_for_the_imported_exam(self):
        source = (Path(self.app.root_path) / "static" / "app.js").read_text(encoding="utf-8")
        start = source.index("async function commitImportPreview")
        end = source.index("\nfunction ", start)
        commit_source = source[start:end]
        self.assertIn("const importedExamKey = result.exam_keys?.[0]", commit_source)
        self.assertIn("state.benchmarkExamKeys = importedExamKey ? [importedExamKey] : [];", commit_source)
        self.assertIn("state.benchmarkInitialized = true;", commit_source)
        self.assertIn('await switchView("benchmark");', commit_source)
        self.assertNotIn('await switchView("overview");', commit_source)

    def test_physics_paired_subject_lines_always_use_chemistry_and_biology(self):
        connection = get_db()
        exam, result_ids = self._prepare_paired_subject_line_case(connection, "物理")
        lines = _derived_subject_lines(connection, exam["exam_key"], "物理")
        for authority, shared in (("chemistry", "politics"), ("biology", "geography")):
            with self.subTest(authority=authority, shared=shared):
                independent_values_differ = False
                for line_type, total_line in (("特控", 500), ("211", 600), ("985", 650)):
                    expected = self._expected_subject_line(connection, result_ids, authority, total_line)
                    independent = self._expected_subject_line(connection, result_ids, shared, total_line)
                    independent_values_differ = independent_values_differ or expected != independent
                    self.assertEqual(lines[authority][line_type], expected)
                    self.assertEqual(lines[shared][line_type], expected)
                    self.assertEqual(lines[authority][line_type] * 2, round(lines[authority][line_type] * 2))
                self.assertTrue(independent_values_differ)

    def test_history_paired_subject_lines_always_use_politics_and_geography(self):
        connection = get_db()
        exam, result_ids = self._prepare_paired_subject_line_case(connection, "历史")
        lines = _derived_subject_lines(connection, exam["exam_key"], "历史")
        for authority, shared in (("politics", "chemistry"), ("geography", "biology")):
            with self.subTest(authority=authority, shared=shared):
                independent_values_differ = False
                for line_type, total_line in (("特控", 500), ("211", 600), ("985", 650)):
                    expected = self._expected_subject_line(connection, result_ids, authority, total_line)
                    independent = self._expected_subject_line(connection, result_ids, shared, total_line)
                    independent_values_differ = independent_values_differ or expected != independent
                    self.assertEqual(lines[authority][line_type], expected)
                    self.assertEqual(lines[shared][line_type], expected)
                    self.assertEqual(lines[authority][line_type] * 2, round(lines[authority][line_type] * 2))
                self.assertTrue(independent_values_differ)

    def test_one_sided_manual_paired_line_is_synchronized(self):
        connection = get_db()
        seed_demo_data(connection)
        exam_key = connection.execute(
            "SELECT exam_key FROM exams ORDER BY exam_date DESC LIMIT 1"
        ).fetchone()["exam_key"]
        response = self.app.test_client().put(
            "/api/settings/score-lines",
            json={
                "exam_key": exam_key,
                "track": "物理",
                "特控": 500,
                "211": 600,
                "985": 650,
                "subject_lines": {
                    "politics": {"特控": 71, "211": 81, "985": 91}
                },
            },
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(
            response.get_json()["subject_lines"]["chemistry"],
            {"特控": 71.0, "211": 81.0, "985": 91.0},
        )
        lines = _derived_subject_lines(connection, exam_key, "物理")
        self.assertEqual(lines["chemistry"], lines["politics"])
        self.assertEqual(lines["chemistry"], {"特控": 71.0, "211": 81.0, "985": 91.0})

    def test_conflicting_manual_paired_lines_are_rejected(self):
        connection = get_db()
        seed_demo_data(connection)
        exam_key = connection.execute(
            "SELECT exam_key FROM exams ORDER BY exam_date DESC LIMIT 1"
        ).fetchone()["exam_key"]
        response = self.app.test_client().put(
            "/api/settings/score-lines",
            json={
                "exam_key": exam_key,
                "track": "历史",
                "特控": 503,
                "211": 603,
                "985": 653,
                "subject_lines": {
                    "politics": {"特控": 70, "211": 80, "985": 90},
                    "chemistry": {"特控": 69, "211": 79, "985": 89},
                },
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("配对科目划线冲突", response.get_json()["error"])

    def test_school_reference_classes_are_created_even_without_results(self):
        workbook_path = Path(self.tempdir.name) / "with_unreferenced_class.xlsx"
        self._make_valid_workbook(workbook_path)
        workbook = load_workbook(workbook_path)
        workbook["学校基础信息表"].append(
            ["重庆", "测试中学", "高2027届", "50", "物理", 40, 8]
        )
        workbook.save(workbook_path)
        connection = get_db()
        preview = preview_workbook(connection, workbook_path, workbook_path.name)
        commit_batch(connection, preview["batch_key"], "reject")
        self.assertIsNotNone(
            connection.execute(
                "SELECT id FROM classes WHERE class_name='50' AND track='物理'"
            ).fetchone()
        )





    def test_comparison_exam_contract_and_selected_previous_exam(self):
        connection = get_db()
        seed_demo_data(connection)
        exams = connection.execute(
            "SELECT id, exam_key, name, exam_date, grade_id FROM exams ORDER BY exam_date, id"
        ).fetchall()
        current = exams[-1]
        selected = comparison_exams(
            connection, current, [current["exam_key"], exams[0]["exam_key"]]
        )
        self.assertEqual([item["exam_key"] for item in selected], [exams[0]["exam_key"], current["exam_key"]])
        with self.assertRaisesRegex(ValueError, "必须包含"):
            comparison_exams(connection, current, [exams[0]["exam_key"]])

        groups = student_groups(
            connection,
            Scope(current["exam_key"], "物理"),
            40,
            [exams[0]["exam_key"], current["exam_key"]],
        )
        self.assertEqual(groups["previous_exam"]["exam_key"], exams[0]["exam_key"])

        history = overview(
            connection,
            Scope(current["exam_key"], "物理"),
            40,
            [exams[0]["exam_key"], current["exam_key"]],
        )["history"]
        for snapshot in history:
            exam_id = connection.execute(
                "SELECT id FROM exams WHERE exam_key=?", (snapshot["exam_key"],)
            ).fetchone()["id"]
            scores = [
                float(row["effective_total"])
                for row in connection.execute(
                    """
                    SELECT er.effective_total
                    FROM exam_results er
                    JOIN student_exam_enrollments se ON se.id=er.enrollment_id
                    WHERE er.exam_id=? AND se.track='物理'
                      AND er.attendance_status='normal' AND er.effective_total IS NOT NULL
                    """,
                    (exam_id,),
                )
            ]
            lines = _line_map(connection, exam_id, "物理")
            self.assertEqual(snapshot["reference_count"], len(scores))
            self.assertEqual(snapshot["critical_count"], sum(lines["特控"] - 40 <= score < lines["特控"] for score in scores))
            self.assertEqual(snapshot["control_count"], sum(score >= lines["特控"] for score in scores))
            self.assertEqual(snapshot["line_211_count"], sum(score >= lines["211"] for score in scores))
            self.assertEqual(snapshot["line_985_count"], sum(score >= lines["985"] for score in scores))

        for index in range(4):
            connection.execute(
                "INSERT INTO exams(exam_key, grade_id, name, exam_date) VALUES (?, ?, ?, ?)",
                (f"limit-{index}", current["grade_id"], f"数量测试{index}", f"2027-0{index+1}-01"),
            )
        connection.commit()
        with self.assertRaisesRegex(ValueError, "最多选择5场"):
            comparison_exams(
                connection,
                current,
                [current["exam_key"], exams[0]["exam_key"], "limit-0", "limit-1", "limit-2", "limit-3"],
            )


    def test_export_options_returns_union_for_multiple_exams(self):
        connection = get_db()
        seed_demo_data(connection)
        exam_keys = [
            row["exam_key"] for row in connection.execute(
                "SELECT exam_key FROM exams ORDER BY exam_date, id"
            ).fetchall()
        ]
        client = self.app.test_client()
        query_string = [("exam_keys", exam_keys[0]), ("exam_keys", exam_keys[-1])]
        response = client.get("/api/export/options", query_string=query_string)
        self.assertEqual(response.status_code, 200, response.get_json())
        payload = response.get_json()
        self.assertGreaterEqual(len(payload["organizations"]), 1)
        self.assertGreaterEqual(len(payload["classes"]), 4)
        self.assertTrue(all("organization_id" in item for item in payload["classes"]))

    def test_successful_batch_list_and_batch_cleanup(self):
        connection = get_db()
        first = Path(self.tempdir.name) / "first.xlsx"
        self._make_valid_workbook(first)
        preview = preview_workbook(connection, first, first.name)
        self.assertEqual(list_batches(connection), [])
        commit_batch(connection, preview["batch_key"], "reject")
        self.assertEqual(len(list_batches(connection)), 1)

        deleted = delete_import_batch(connection, preview["batch_key"])
        self.assertEqual(deleted["deleted_results"], 1)
        self.assertEqual(connection.execute("SELECT COUNT(*) AS n FROM exam_results").fetchone()["n"], 0)
        self.assertEqual(list_batches(connection), [])

        second = Path(self.tempdir.name) / "second.xlsx"
        self._make_valid_workbook(second)
        preview = preview_workbook(connection, second, second.name)
        commit_batch(connection, preview["batch_key"], "reject")
        cleared = clear_import_cache(connection)
        self.assertEqual(cleared["deleted_batches"], 1)
        self.assertEqual(cleared["deleted_results"], 1)
        self.assertEqual(connection.execute("SELECT COUNT(*) AS n FROM import_batches").fetchone()["n"], 0)

    def test_priority1_migration_merges_safe_legacy_student_duplicates(self):
        connection = get_db()
        workbook_path = Path(self.tempdir.name) / "deduplicate.xlsx"
        self._make_valid_workbook(workbook_path)
        preview = preview_workbook(connection, workbook_path, workbook_path.name)
        result = commit_batch(connection, preview["batch_key"], "reject")
        exam = connection.execute(
            "SELECT id, exam_key, grade_id FROM exams WHERE exam_key=?", (result["exam_keys"][0],)
        ).fetchone()
        enrollment = connection.execute(
            "SELECT organization_id, class_id FROM student_exam_enrollments WHERE exam_id=?",
            (exam["id"],),
        ).fetchone()
        duplicate = connection.execute(
            "INSERT INTO students(student_no, display_student_no, identity_key, student_type, name) VALUES ('legacy-duplicate', 'S001', 'legacy-repeat-S001', '普通生', '测试学生')"
        )
        duplicate_student_id = duplicate.lastrowid
        second_exam = connection.execute(
            """INSERT INTO exams(exam_key, grade_id, name, exam_date)
               VALUES ('second-exam', ?, '第二次考试', '2026-06-12')""",
            (exam["grade_id"],),
        )
        second_exam_id = second_exam.lastrowid
        second_enrollment = connection.execute(
            """
            INSERT INTO student_exam_enrollments(
              student_id, exam_id, organization_id, grade_id, class_id, track,
              elective_subjects_json, student_type, source
            ) VALUES (?, ?, ?, ?, ?, '物理', '[]', '普通生', 'import')
            """,
            (duplicate_student_id, second_exam_id, enrollment["organization_id"], exam["grade_id"], enrollment["class_id"]),
        )
        connection.execute(
            """
            INSERT INTO exam_results(
              exam_id, student_id, enrollment_id, effective_total, total_source,
              attendance_status
            ) VALUES (?, ?, ?, 530, 'scaled', 'normal')
            """,
            (second_exam_id, duplicate_student_id, second_enrollment.lastrowid),
        )
        connection.execute(
            "DELETE FROM operation_logs WHERE target_type='data_migration' AND target_id=?",
            (CORE_RULE_VERSION,),
        )
        connection.commit()

        migrated = run_priority1_data_migration(connection)
        self.assertEqual(migrated["merged_students"], 1)
        students = connection.execute(
            "SELECT id, identity_key FROM students WHERE display_student_no='S001' AND name='测试学生'"
        ).fetchall()
        self.assertEqual(len(students), 1)
        self.assertEqual(students[0]["identity_key"], "regular|s001|测试学生")
        result_student_ids = {
            row["student_id"] for row in connection.execute(
                "SELECT student_id FROM exam_results ORDER BY exam_id"
            )
        }
        self.assertEqual(result_student_ids, {students[0]["id"]})
        self.assertEqual(len(student_detail(connection, students[0]["identity_key"])["history"]), 2)

    def test_priority1_migration_stops_on_same_exam_student_conflict(self):
        connection = get_db()
        workbook_path = Path(self.tempdir.name) / "conflict.xlsx"
        self._make_valid_workbook(workbook_path)
        preview = preview_workbook(connection, workbook_path, workbook_path.name)
        result = commit_batch(connection, preview["batch_key"], "reject")
        exam = connection.execute(
            "SELECT id, grade_id FROM exams WHERE exam_key=?", (result["exam_keys"][0],)
        ).fetchone()
        enrollment = connection.execute(
            "SELECT organization_id, class_id FROM student_exam_enrollments WHERE exam_id=?",
            (exam["id"],),
        ).fetchone()
        duplicate_id = connection.execute(
            "INSERT INTO students(student_no, display_student_no, identity_key, student_type, name) VALUES ('legacy-conflict', 'S001', 'legacy-conflict-S001', '普通生', '测试学生')"
        ).lastrowid
        enrollment_id = connection.execute(
            """INSERT INTO student_exam_enrollments(
                 student_id, exam_id, organization_id, grade_id, class_id, track,
                 elective_subjects_json, student_type, source
               ) VALUES (?, ?, ?, ?, ?, '物理', '[]', '普通生', 'import')""",
            (duplicate_id, exam["id"], enrollment["organization_id"], exam["grade_id"], enrollment["class_id"]),
        ).lastrowid
        connection.execute(
            """INSERT INTO exam_results(
                 exam_id, student_id, enrollment_id, effective_total, total_source, attendance_status
               ) VALUES (?, ?, ?, 520, 'scaled', 'normal')""",
            (exam["id"], duplicate_id, enrollment_id),
        )
        connection.execute(
            "DELETE FROM operation_logs WHERE target_type='data_migration' AND target_id=?",
            (CORE_RULE_VERSION,),
        )
        connection.commit()
        with self.assertRaisesRegex(RuntimeError, "同场考试冲突"):
            run_priority1_data_migration(connection)
        self.assertEqual(
            connection.execute(
                "SELECT COUNT(*) n FROM students WHERE display_student_no='S001' AND name='测试学生'"
            ).fetchone()["n"],
            2,
        )
        self.assertEqual(connection.execute("SELECT COUNT(*) n FROM exam_results").fetchone()["n"], 2)

    def test_two_file_import_api_download_commit_and_delete(self):
        source = Path(self.tempdir.name) / "api_source.xlsx"
        self._make_valid_workbook(source)
        workbook = load_workbook(source)
        extra_sheet = workbook.create_sheet("其他学生数据", 4)
        extra_sheet.append(["年级", "考试名称", "考试时间", "学生类型", "班级", "姓名", "物理", "历史", "总分"])
        workbook.save(source)
        source_bytes = source.read_bytes()
        client = self.app.test_client()

        response = client.post(
            "/api/import/prepare",
            data={
                "reference_file": (BytesIO(source_bytes), "基础信息.xlsx"),
                "score_file": (BytesIO(source_bytes), "成绩总表.xlsx"),
                "grade": "高2028届",
                "exam_name": "高二下期末考试",
                "exam_month": "2027年06月",
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        preview = response.get_json()
        self.assertEqual(preview["summary"]["preparation"]["matched_rows"], 1)
        download = client.get(f"/api/import/{preview['batch_key']}/workbook")
        self.assertEqual(download.status_code, 200)
        self.assertGreater(len(download.data), 1000)
        download.close()

        committed = client.post(
            f"/api/import/{preview['batch_key']}/commit",
            json={"mode": "reject", "warnings_confirmed": True},
        )
        self.assertEqual(committed.status_code, 200, committed.get_json())
        batches = client.get("/api/import/batches").get_json()
        self.assertEqual(batches["count"], 1)
        deleted = client.delete(f"/api/import/{preview['batch_key']}")
        self.assertEqual(deleted.status_code, 200, deleted.get_json())
        self.assertEqual(deleted.get_json()["deleted_results"], 1)

    def test_clear_demo_preserves_imported_data_and_upload_batches(self):
        connection = get_db()
        seed_demo_data(connection)
        workbook_path = Path(self.tempdir.name) / "keep_import.xlsx"
        self._make_valid_workbook(workbook_path)
        preview = preview_workbook(connection, workbook_path, workbook_path.name)
        result = commit_batch(connection, preview["batch_key"], "reject")

        counts = clear_demo_data(connection)

        self.assertGreater(counts["results"], 0)
        self.assertEqual(connection.execute("SELECT COUNT(*) AS n FROM exams").fetchone()["n"], 1)
        self.assertEqual(connection.execute("SELECT COUNT(*) AS n FROM exam_results").fetchone()["n"], 1)
        self.assertEqual(connection.execute("SELECT COUNT(*) AS n FROM students WHERE student_no LIKE 'DEMO%'").fetchone()["n"], 0)
        batch = connection.execute(
            "SELECT status FROM import_batches WHERE batch_key=?", (result["batch_key"],)
        ).fetchone()
        self.assertEqual(batch["status"], "committed")

    def test_formula_student_number_uses_name_and_class_fallback(self):
        workbook_path = Path(self.tempdir.name) / "formula_import.xlsx"
        self._make_valid_workbook(workbook_path)
        workbook = load_workbook(workbook_path)
        workbook["物理方向学生成绩"]["D2"] = '=XLOOKUP(F2,学生信息!A:A,学生信息!B:B)'
        workbook.save(workbook_path)
        payload = parse_workbook(workbook_path, get_db())
        errors = [issue for issue in payload.issues if issue.severity == "error"]
        self.assertEqual(errors, [])
        self.assertEqual(payload.scores[0]["student_no"], "S001")
        self.assertEqual(payload.scores[0]["identity_key"], "regular|s001|测试学生")
        self.assertTrue(
            any(issue.issue_code == "student_identity_fallback" for issue in payload.issues)
        )

    def test_template_metrics_use_left_closed_bands_and_half_point_subject_lines(self):
        workbook_path = Path(self.tempdir.name) / "metric_rules.xlsx"
        self._make_valid_workbook(workbook_path)
        connection = get_db()
        preview = preview_workbook(connection, workbook_path, workbook_path.name)
        result = commit_batch(connection, preview["batch_key"], "reject")
        exam = connection.execute("SELECT id, exam_key FROM exams").fetchone()
        connection.execute(
            "UPDATE exam_results SET effective_total=500, scaled_total=500"
        )
        connection.execute(
            "UPDATE score_lines SET line_value=CASE line_type WHEN '特控' THEN 500 WHEN '211' THEN 600 ELSE 650 END WHERE exam_id=? AND track='物理'",
            (exam["id"],),
        )
        connection.commit()
        scope = Scope(exam["exam_key"], "物理")

        bands = score_bands(connection, scope)
        line_counts = {item["label"]: item["count"] for item in bands["line_bands"]}
        fixed_counts = {item["label"]: item["count"] for item in bands["fixed_bands"]}
        self.assertEqual(line_counts["460-500"], 0)
        self.assertEqual(line_counts["500-600"], 1)
        self.assertEqual(fixed_counts["450—500"], 0)
        self.assertEqual(fixed_counts["500—550"], 1)

        subjects = subject_metrics(connection, scope)
        for item in subjects["subjects"]:
            for line in item["lines"].values():
                if line is not None:
                    self.assertEqual(line * 2, round(line * 2))
        classes = class_metrics(connection, scope)
        self.assertEqual(classes["classes"][0]["control_count"], 1)
        self.assertEqual(classes["classes"][0]["potential_count"], 1)
        matrix = student_group_matrix(connection, scope)
        self.assertEqual(matrix["students"][0]["totals"][-1]["level"], "line-control")
        self.assertEqual(result["exam_keys"], [exam["exam_key"]])

    def test_same_student_number_with_different_names_in_one_exam_is_blocked(self):
        workbook_path = Path(self.tempdir.name) / "same_exam_number_conflict.xlsx"
        self._make_valid_workbook(workbook_path)
        workbook = load_workbook(workbook_path)
        students = workbook["学生信息"]
        students.append(["另一学生", "S001", "高2027届", "2", "测试中学", "物化生"])
        physics = workbook["物理方向学生成绩"]
        row = list(physics.iter_rows(min_row=2, max_row=2, max_col=48, values_only=True))[0]
        row = list(row); row[4] = "2"; row[5] = "另一学生"; physics.append(row)
        workbook.save(workbook_path)
        preview = preview_workbook(get_db(), workbook_path, workbook_path.name)
        conflicts = [
            issue for issue in preview["issues"]
            if issue["issue_code"] == "exam_student_number_name_conflict"
        ]
        self.assertEqual(len(conflicts), 2)
        with self.assertRaisesRegex(ValueError, "阻止入库"):
            commit_batch(get_db(), preview["batch_key"], "reject")

    def test_same_number_and_name_across_class_and_track_keeps_student_id(self):
        connection = get_db()
        first = Path(self.tempdir.name) / "identity_first.xlsx"
        second = Path(self.tempdir.name) / "identity_second.xlsx"
        self._make_valid_workbook(first)
        self._make_valid_workbook(second)
        first_preview = preview_workbook(connection, first, first.name)
        first_result = commit_batch(connection, first_preview["batch_key"], "reject")
        self._customize_valid_workbook(
            second,
            exam_name="第二次考试",
            exam_date=date(2026, 6, 12),
            school="测试中学",
            class_name="2",
            track="历史",
        )
        second_preview = preview_workbook(connection, second, second.name)
        self.assertEqual(second_preview["error_count"], 0, second_preview["issues"])
        second_result = commit_batch(connection, second_preview["batch_key"], "reject")
        students = connection.execute(
            "SELECT id, identity_key FROM students WHERE display_student_no='S001' AND name='测试学生'"
        ).fetchall()
        self.assertEqual(len(students), 1)
        enrollments = connection.execute(
            """SELECT se.student_id, se.track, o.name organization_name, c.class_name
               FROM student_exam_enrollments se
               JOIN organizations o ON o.id=se.organization_id
               JOIN classes c ON c.id=se.class_id
               WHERE se.student_id=? ORDER BY se.exam_id""",
            (students[0]["id"],),
        ).fetchall()
        self.assertEqual(
            [(row["track"], row["organization_name"], row["class_name"]) for row in enrollments],
            [("物理", "测试中学", "1"), ("历史", "测试中学", "2")],
        )
        detail = student_detail(connection, students[0]["identity_key"])
        self.assertEqual(len(detail["history"]), 2)
        self.assertEqual(
            {item["exam_key"] for item in detail["history"]},
            {first_result["exam_keys"][0], second_result["exam_keys"][0]},
        )

    def test_same_number_different_names_across_exams_are_distinct_students(self):
        connection = get_db()
        first = Path(self.tempdir.name) / "reused_number_first.xlsx"
        second = Path(self.tempdir.name) / "reused_number_second.xlsx"
        self._make_valid_workbook(first)
        self._make_valid_workbook(second)
        first_preview = preview_workbook(connection, first, first.name)
        commit_batch(connection, first_preview["batch_key"], "reject")
        self._customize_valid_workbook(
            second,
            exam_name="第二次考试",
            exam_date=date(2026, 6, 12),
            student_name="另一学生",
        )
        second_preview = preview_workbook(connection, second, second.name)
        self.assertEqual(second_preview["error_count"], 0, second_preview["issues"])
        commit_batch(connection, second_preview["batch_key"], "reject")
        identities = connection.execute(
            "SELECT identity_key FROM students WHERE display_student_no='S001' ORDER BY id"
        ).fetchall()
        self.assertEqual(
            [row["identity_key"] for row in identities],
            ["regular|s001|测试学生", "regular|s001|另一学生"],
        )

    def test_different_numbers_with_same_name_are_distinct_students(self):
        workbook_path = Path(self.tempdir.name) / "same_name_different_numbers.xlsx"
        self._make_valid_workbook(workbook_path)
        workbook = load_workbook(workbook_path)
        workbook["学生信息"].append(["测试学生", "S002", "高2027届", "1", "测试中学", "物化生"])
        physics = workbook["物理方向学生成绩"]
        row = list(physics.iter_rows(min_row=2, max_row=2, max_col=48, values_only=True))[0]
        row = list(row); row[3] = "S002"; physics.append(row)
        workbook.save(workbook_path)
        preview = preview_workbook(get_db(), workbook_path, workbook_path.name)
        self.assertEqual(preview["error_count"], 0, preview["issues"])
        commit_batch(get_db(), preview["batch_key"], "reject")
        identities = get_db().execute(
            "SELECT identity_key FROM students WHERE name='测试学生' ORDER BY identity_key"
        ).fetchall()
        self.assertEqual(
            [row["identity_key"] for row in identities],
            ["regular|s001|测试学生", "regular|s002|测试学生"],
        )

    def test_historical_class_direction_cannot_be_overwritten_by_base_upload(self):
        connection = get_db()
        workbook_path = Path(self.tempdir.name) / "historical_class.xlsx"
        self._make_valid_workbook(workbook_path)
        preview = preview_workbook(connection, workbook_path, workbook_path.name)
        commit_batch(connection, preview["batch_key"], "reject")
        enrollment_before = connection.execute(
            "SELECT class_id, track FROM student_exam_enrollments"
        ).fetchone()
        workbook = load_workbook(workbook_path)
        workbook["学校基础信息表"].cell(2, 5, "历史")
        workbook.save(workbook_path)
        with self.assertRaisesRegex(ValueError, "历史考试不会被自动覆盖"):
            save_base_information(connection, parse_base_information_workbook(workbook_path))
        enrollment_after = connection.execute(
            "SELECT class_id, track FROM student_exam_enrollments"
        ).fetchone()
        self.assertEqual(tuple(enrollment_after), tuple(enrollment_before))
        self.assertEqual(enrollment_after["track"], "物理")



    def test_orphan_grade_cleanup_keeps_every_referenced_grade(self):
        connection = get_db()
        connection.execute("INSERT INTO grades(name) VALUES ('孤立年级')")
        connection.execute("INSERT INTO grades(name) VALUES ('基础信息保留年级')")
        connection.execute(
            """INSERT INTO base_school_classes(
                 school_name, grade_name, class_name, track
               ) VALUES ('测试中学', '基础信息保留年级', '1', '物理')"""
        )
        connection.execute(
            "DELETE FROM operation_logs WHERE target_type='data_migration' AND target_id=?",
            (CORE_RULE_VERSION,),
        )
        connection.commit()
        result = run_priority1_data_migration(connection)
        self.assertEqual(result["orphan_grades"], 1)
        self.assertIsNone(
            connection.execute("SELECT id FROM grades WHERE name='孤立年级'").fetchone()
        )
        self.assertIsNotNone(
            connection.execute("SELECT id FROM grades WHERE name='基础信息保留年级'").fetchone()
        )
        log = connection.execute(
            """SELECT summary FROM operation_logs
               WHERE action='orphan_grades_cleanup' AND target_id=?""",
            (CORE_RULE_VERSION,),
        ).fetchone()
        self.assertIn('"deleted_grades": 1', log["summary"])

    def test_class_and_overall_target_rates_keep_distinct_formulas(self):
        connection = get_db()
        workbook_path = Path(self.tempdir.name) / "target_formulas.xlsx"
        self._make_valid_workbook(workbook_path)
        preview = preview_workbook(connection, workbook_path, workbook_path.name)
        result = commit_batch(connection, preview["batch_key"], "reject")
        connection.execute("UPDATE class_targets SET target_count=2")
        connection.commit()
        class_result = class_metrics(connection, Scope(result["exam_keys"][0], "物理"), target_name="特控")
        overall_result = overview(connection, Scope(result["exam_keys"][0], "物理"), target_name="特控")
        self.assertEqual(class_result["classes"][0]["target_completion_rate"], 0.5)
        self.assertEqual(overall_result["targets"]["rate"], 0)

    def test_target_selection_changes_only_target_metrics(self):
        connection = get_db()
        seed_demo_data(connection)
        exam = connection.execute(
            "SELECT id, exam_key FROM exams ORDER BY exam_date DESC, id DESC LIMIT 1"
        ).fetchone()
        class_ids = [row["class_id"] for row in connection.execute(
            """SELECT DISTINCT se.class_id FROM student_exam_enrollments se
               JOIN exam_results er ON er.enrollment_id=se.id
               WHERE se.exam_id=? AND se.track='物理' AND er.attendance_status='normal'""",
            (exam["id"],),
        )]
        for class_id in class_ids:
            connection.execute(
                "INSERT INTO class_stage_targets(class_id,target_name,target_order,target_count,source) VALUES (?, '一诊', 0, 2, 'test')",
                (class_id,),
            )
            connection.execute(
                "INSERT INTO class_stage_targets(class_id,target_name,target_order,target_count,source) VALUES (?, '二诊', 1, 4, 'test')",
                (class_id,),
            )
        connection.commit()
        first = overview(connection, Scope(exam["exam_key"], "物理"), target_name="一诊")
        second = overview(connection, Scope(exam["exam_key"], "物理"), target_name="二诊")
        self.assertEqual(first["average"], second["average"])
        self.assertNotEqual(
            [item["target"] for item in first["targets"]["details"]],
            [item["target"] for item in second["targets"]["details"]],
        )
        class_first = class_metrics(connection, Scope(exam["exam_key"], "物理"), target_name="一诊")
        class_second = class_metrics(connection, Scope(exam["exam_key"], "物理"), target_name="二诊")
        self.assertEqual(class_first["classes"][0]["average"], class_second["classes"][0]["average"])
        self.assertEqual(class_first["classes"][0]["target"], 2)
        self.assertEqual(class_second["classes"][0]["target"], 4)

    def test_student_search_queries_full_database_before_limit(self):
        connection = get_db()
        seed_demo_data(connection)
        exam = connection.execute(
            "SELECT id, exam_key, grade_id FROM exams ORDER BY exam_date DESC, id DESC LIMIT 1"
        ).fetchone()
        organization_id = connection.execute("INSERT INTO organizations(name) VALUES ('深层测试中学')").lastrowid
        class_id = connection.execute(
            "INSERT INTO classes(organization_id,grade_id,class_name,track) VALUES (?,?,'701','物理')",
            (organization_id, exam["grade_id"]),
        ).lastrowid
        for index in range(701):
            internal = f"search-internal-{index:04d}"
            display = f"SEARCH{index:04d}"
            name = "数据库深层学生" if index == 700 else f"搜索占位{index:04d}"
            student_id = connection.execute(
                "INSERT INTO students(student_no,display_student_no,identity_key,name,student_type) VALUES (?,?,?,?, '普通生')",
                (internal, display, internal, name),
            ).lastrowid
            enrollment_id = connection.execute(
                """INSERT INTO student_exam_enrollments(
                     student_id,exam_id,organization_id,grade_id,class_id,track,elective_subjects_json,student_type
                   ) VALUES (?,?,?,?,?,'物理','[]','普通生')""",
                (student_id, exam["id"], organization_id, exam["grade_id"], class_id),
            ).lastrowid
            score = 700 - index * 0.1
            connection.execute(
                """INSERT INTO exam_results(
                     exam_id,student_id,enrollment_id,raw_total,scaled_total,effective_total,total_source
                   ) VALUES (?,?,?,?,?,?,'scaled')""",
                (exam["id"], student_id, enrollment_id, score, score, score),
            )
        connection.commit()
        scope = Scope(exam["exam_key"], "物理")
        self.assertEqual(len(student_list(connection, scope, 10)), 10)
        for keyword in ("数据库深层学生", "SEARCH0700", "深层测试中学", "701"):
            matches = student_list(connection, scope, 50, keyword)
            self.assertTrue(matches, keyword)
            self.assertLessEqual(len(matches), 50)
        response = self.app.test_client().get(
            "/api/students",
            query_string={"exam_key": exam["exam_key"], "track": "物理", "q": "SEARCH0700"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["items"][0]["name"], "数据库深层学生")

    def test_class_metrics_lists_all_participated_exams_outside_selected_history(self):
        connection = get_db()
        seed_demo_data(connection)
        exams = connection.execute("SELECT exam_key FROM exams ORDER BY exam_date, id").fetchall()
        current_key = exams[-1]["exam_key"]
        class_id = connection.execute(
            """SELECT se.class_id FROM student_exam_enrollments se
               JOIN exams e ON e.id=se.exam_id
               WHERE e.exam_key=? AND se.track='物理' LIMIT 1""",
            (current_key,),
        ).fetchone()["class_id"]
        result = class_metrics(
            connection, Scope(current_key, "物理", class_id=class_id), exam_keys=[current_key]
        )
        participated = result["classes"][0]["participated_exams"]
        self.assertGreater(len(participated), 1)
        self.assertTrue(any(item["exam_key"] != current_key for item in participated))

    def test_nonstandard_score_headers_are_normalized(self):
        path = Path(self.tempdir.name) / "headers.xlsx"
        self._make_valid_workbook(path)
        workbook = load_workbook(path)
        sheet = workbook["物理方向学生成绩"]
        for column, value in {9:"YWM", 21:"化学M", 22:"化学N", 23:"化学N排名", 24:"HXM"}.items():
            sheet.cell(1, column, value)
        workbook.save(path)
        payload = parse_workbook(path, get_db())
        mapped = {(item["column"], item["normalized"]) for item in payload.header_mappings}
        self.assertTrue({(21,"化学等级"),(22,"化学赋分"),(23,"化学赋分排名")}.issubset(mapped))
        self.assertFalse(any(item["column"] in (9,24) for item in payload.header_mappings))


    def test_uploaded_external_average_is_not_imported(self):
        path = Path(self.tempdir.name) / "external_average.xlsx"
        self._make_valid_workbook(path)
        workbook = load_workbook(path)
        workbook["额外参考均分"].append(["高2027届", "高二月考", date(2026,5,12), "物理", 999])
        workbook.save(path)
        payload = parse_workbook(path, get_db())
        self.assertFalse(hasattr(payload, "benchmark_averages"))
        preview = preview_workbook(get_db(), path, path.name)
        result = commit_batch(get_db(), preview["batch_key"])
        data = benchmark_metrics(get_db(), result["exam_keys"][0])
        self.assertEqual(data["tracks"]["物理"]["grade_averages"]["total"], 540)
        self.assertEqual(get_db().execute("SELECT COUNT(*) FROM sqlite_master WHERE name='benchmark_averages'").fetchone()[0], 0)



    def test_regular_identity_and_score_blocks_can_be_reordered(self):
        workbook_path = Path(self.tempdir.name) / "reordered_regular.xlsx"
        self._make_valid_workbook(workbook_path)
        workbook = load_workbook(workbook_path)
        sheet = workbook["物理方向学生成绩"]
        headers = [cell.value for cell in sheet[1]]
        values = [cell.value for cell in sheet[2]]
        order = [6, 1, 3, 5, 4, 2, *range(10, 13), *range(43, 49), *range(7, 10), *range(13, 43)]
        self.assertEqual(sorted(order), list(range(1, 49)))
        sheet.delete_rows(1, sheet.max_row)
        sheet.append([headers[index - 1] for index in order])
        sheet.append([values[index - 1] for index in order])
        workbook.save(workbook_path)
        payload = parse_workbook(workbook_path, get_db())
        errors = [issue for issue in payload.issues if issue.severity == "error"]
        self.assertEqual(errors, [])
        self.assertEqual(payload.scores[0]["student_no"], "S001")
        self.assertEqual(payload.scores[0]["effective_total"], 540)


    def test_class_exports_preserve_borders(self):
        connection = get_db()
        seed_demo_data(connection)
        exam_key = connection.execute("SELECT exam_key FROM exams ORDER BY exam_date DESC LIMIT 1").fetchone()["exam_key"]
        exam = connection.execute(
            "SELECT grade_id FROM exams WHERE exam_key=?", (exam_key,)
        ).fetchone()
        organization_id = connection.execute(
            "SELECT id FROM organizations WHERE name='测试中学'"
        ).fetchone()["id"]
        connection.execute(
            """
            INSERT INTO classes(organization_id, grade_id, class_name, track, planned_size, control_target)
            VALUES (?, ?, '99', '物理', 30, 15)
            """,
            (organization_id, exam["grade_id"]),
        )
        for class_row in connection.execute(
            "SELECT id FROM classes WHERE grade_id=?",
            (exam["grade_id"],),
        ):
            for target_order, (target_name, target_count) in enumerate(
                (("一诊", 12), ("二诊", 14), ("高考", 16))
            ):
                connection.execute(
                    """INSERT INTO class_stage_targets(
                           class_id,target_name,target_order,target_count,source
                       ) VALUES (?,?,?,?, 'test')""",
                    (class_row["id"], target_name, target_order, target_count),
                )
        connection.commit()
        archive = export_class_zip(connection, exam_key, self.app.config["EXPORT_FOLDER"])
        self.assertTrue(archive.exists())
        with zipfile.ZipFile(archive) as class_archive:
            class_name = next(name for name in class_archive.namelist() if name.endswith(".xlsx"))
            class_workbook = load_workbook(BytesIO(class_archive.read(class_name)))
        class_sheet = class_workbook[class_workbook.sheetnames[0]]
        for row in class_sheet.iter_rows(min_row=1, max_row=class_sheet.max_row, min_col=1, max_col=class_sheet.max_column):
            for cell in row:
                self.assertEqual(cell.border.left.style, "thin")
                self.assertTrue(cell.border.left.color.rgb.endswith("000000"))
        class_workbook.close()

    def test_multi_exam_class_exports_reuse_single_exam_outputs(self):
        connection = get_db()
        seed_demo_data(connection)
        exams = connection.execute(
            "SELECT exam_key, name FROM exams ORDER BY exam_date, id"
        ).fetchall()
        selected_keys = [exams[0]["exam_key"], exams[-1]["exam_key"]]
        class_id = connection.execute(
            "SELECT id FROM classes WHERE class_name='01' LIMIT 1"
        ).fetchone()["id"]
        class_zip = export_class_zip(
            connection,
            selected_keys[0],
            self.app.config["EXPORT_FOLDER"],
            exam_keys=selected_keys,
            class_ids=[class_id],
        )
        with zipfile.ZipFile(class_zip) as archive:
            names = [name for name in archive.namelist() if name.endswith(".xlsx")]
        self.assertEqual(len(names), 2)
        self.assertEqual(len({name.split("/", 1)[0] for name in names}), 2)

    def test_export_page_uses_independent_filters_and_student_scope(self):
        source = (Path(__file__).parents[1] / "app" / "static" / "app.js").read_text(encoding="utf-8")
        self.assertIn("exportFilters: {", source)
        for key in ("classes",):
            self.assertIn(f"{key}:", source)
        self.assertIn('["import","benchmark","export"].includes(state.view)', source)
        self.assertIn('query({limit:"10"})', source)
        self.assertIn('query({q:keyword,limit:keyword?"50":"10"})', source)
        self.assertIn('selectedTargetStorageKey = "student-analysis-selected-target"', source)
        self.assertIn('breakOnMissing=false', source)
        self.assertIn('comparisonExamsStorageKey = "student-analysis-comparison-exams"', source)
        self.assertIn('maxSelections:5,lockedValues:[state.filters.exam_key]', source)
        self.assertIn('groupTransitionPageSize: 20', source)
        self.assertIn('targetPageSize: 10', source)
        self.assertIn('["chinese","math","english","physics","history","chemistry","biology","politics","geography"]', source)
        self.assertIn('class="benchmark-section benchmark-collapsible"', source)
        self.assertIn('class="matrix-table band-matrix"', source)
        self.assertIn('highlightExamKey:state.filters.exam_key', source)
        self.assertNotIn("benchmarkSearch", source)
        self.assertNotIn("slice(0, 3)", source)




    def test_missing_history_elective_line_uses_paired_subject(self):
        connection = get_db()
        seed_demo_data(connection)
        exam_key = connection.execute(
            "SELECT exam_key FROM exams ORDER BY exam_date DESC LIMIT 1"
        ).fetchone()["exam_key"]
        exam_id = connection.execute(
            "SELECT id FROM exams WHERE exam_key=?", (exam_key,)
        ).fetchone()["id"]
        connection.execute(
            """DELETE FROM subject_scores WHERE subject_code='chemistry' AND exam_result_id IN (
                 SELECT er.id FROM exam_results er
                 JOIN student_exam_enrollments se ON se.id=er.enrollment_id
                 WHERE er.exam_id=? AND se.track='历史'
               )""",
            (exam_id,),
        )
        connection.commit()
        history = benchmark_metrics(connection, exam_key)["tracks"]["历史"]
        lines = {item["code"]: item["lines"] for item in history["subject_lines"]}
        self.assertIsNotNone(lines["politics"]["特控"])
        self.assertEqual(lines["chemistry"]["特控"], lines["politics"]["特控"])


    def _customize_valid_workbook(
        self,
        path: Path,
        *,
        exam_name: str,
        exam_date: date,
        student_no: str = "S001",
        student_name: str = "测试学生",
        school: str = "测试中学",
        class_name: str = "1",
        track: str = "物理",
    ) -> None:
        workbook = load_workbook(path)
        school_sheet = workbook["学校基础信息表"]
        school_sheet.cell(2, 2, school)
        school_sheet.cell(2, 4, class_name)
        school_sheet.cell(2, 5, track)
        student_sheet = workbook["学生信息"]
        student_sheet.cell(2, 1, student_name)
        student_sheet.cell(2, 2, student_no)
        student_sheet.cell(2, 4, class_name)
        student_sheet.cell(2, 5, school)
        student_sheet.cell(2, 6, "物化生" if track == "物理" else "史化生")
        physics = workbook["物理方向学生成绩"]
        history = workbook["历史方向学生成绩"]
        row = list(physics.iter_rows(min_row=2, max_row=2, max_col=48, values_only=True))[0]
        if physics.max_row > 1:
            physics.delete_rows(2, physics.max_row - 1)
        if history.max_row > 1:
            history.delete_rows(2, history.max_row - 1)
        row = list(row)
        row[:6] = ["高2027届", exam_name, exam_date, student_no, class_name, student_name]
        target = physics if track == "物理" else history
        target.cell(1, 16, track)
        target.append(row)
        workbook.save(path)

    def _make_base_grades_workbook(self, path: Path, grades: list[str]) -> None:
        workbook = Workbook()
        school = workbook.active
        school.title = "学校基础信息表"
        school.append(["省份", "学校", "年级", "班级", "方向", "班级人数", "特控目标"])
        students = workbook.create_sheet("学生信息")
        students.append(["姓名", "学号", "年级", "班级", "学校", "选考标签"])
        for index, grade in enumerate(grades, start=1):
            school.append(["重庆", "跨年级测试中学", grade, "1", "物理", 1, 1])
            students.append([f"测试学生{index}", f"G{index:03d}", grade, "1", "跨年级测试中学", "物化生"])
        workbook.save(path)

    def _make_valid_workbook(self, path: Path) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)
        school = workbook.create_sheet("学校基础信息表")
        school.append(["省份", "学校", "年级", "班级", "方向", "班级人数", "特控目标"])
        school.append(["重庆", "测试中学", "高2027届", "1", "物理", 1, 1])
        students = workbook.create_sheet("学生信息")
        students.append(["姓名", "学号", "年级", "班级", "学校", "选考标签"])
        students.append(["测试学生", "S001", "高2027届", "1", "测试中学", "物化生"])
        headers = [
            "年级", "考试名称", "考试时间", "学号", "班级", "姓名",
            "语文", "排名", "外部参考对比排名",
            "数学", "排名", "外部参考对比排名",
            "英语", "排名", "外部参考对比排名",
            "物理", "排名", "外部参考对比排名",
            "化学", "排名", "化学等级", "化学赋分", "化学赋分排名", "赋分外部参考对比排名",
            "生物", "排名", "生物等级", "生物赋分", "生物赋分排名", "赋分外部参考对比排名",
            "政治", "排名", "政治等级", "政治赋分", "政治赋分排名", "赋分外部参考对比排名",
            "地理", "排名", "地理等级", "地理赋分", "地理赋分排名", "赋分外部参考对比排名",
            "总分", "排名", "原始分外部参考对比排名", "总分赋分", "总分赋分排名", "赋分外部参考对比排名",
        ]
        physics = workbook.create_sheet("物理方向学生成绩")
        physics.append(headers)
        row = [None] * 48
        row[:6] = ["高2027届", "高二月考", date(2026, 5, 12), "S001", "1", "测试学生"]
        for column, value in {7: 100, 10: 100, 13: 100, 16: 80, 19: 60, 22: 80, 25: 60, 28: 80, 43: 500, 46: 540}.items():
            row[column - 1] = value
        physics.append(row)
        history = workbook.create_sheet("历史方向学生成绩")
        history.append(headers)
        benchmark = workbook.create_sheet("额外参考均分")
        benchmark.append(["额外参考数据"])
        benchmark.append(["年级", "考试名称", "考试时间", "方向", "总分平均分（赋）"])
        lines = workbook.create_sheet("分数线表")
        for _ in range(15):
            lines.append([None] * 10)
        lines["B4"] = 500
        lines["B5"] = 570
        lines["B6"] = 630
        lines["B8"] = 505
        lines["B9"] = 575
        lines["B10"] = 635
        workbook.save(path)


if __name__ == "__main__":
    unittest.main()
