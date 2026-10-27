from __future__ import annotations

import copy
import re
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.config import SUBJECTS
from app.services.analytics import grade_statistics


SUBJECT_ORDER = {
    "物理": ["chinese", "math", "english", "physics", "chemistry", "biology", "politics", "geography"],
    "历史": ["chinese", "math", "english", "history", "politics", "geography", "biology", "chemistry"],
}


SCORE_LEVEL_FILLS = {
    "line-critical": "F6C8CE",
    "line-control": "C8E5B4",
    "line-211": "B6C7EC",
    "line-985": "F9E3A1",
}
NAVY = "173B63"
PALE_BLUE = "EAF2F8"
PALE_TEAL = "E8F5F2"
PALE_GOLD = "FFF4D6"
GRID = "D9E2EC"
WHITE = "FFFFFF"
THIN_BORDER = Border(
    left=Side(style="thin", color=GRID),
    right=Side(style="thin", color=GRID),
    top=Side(style="thin", color=GRID),
    bottom=Side(style="thin", color=GRID),
)
BLACK_THIN_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)


def _semantic_number_format(label: str, row_label: str = "") -> str:
    normalized = str(label or "")
    row_normalized = str(row_label or "")
    rate_tokens = ("率", "占比", "比值", "班级/年级", "班级/年级")
    rank_tokens = ("排名", "位次")
    count_tokens = ("人数", "班级数", "参考")
    if any(token in row_normalized or token in normalized for token in rank_tokens):
        return "0"
    if any(token in row_normalized for token in rate_tokens):
        return "0.00%"
    if any(token in row_normalized for token in count_tokens):
        return "0"
    if any(token in normalized for token in rate_tokens):
        return "0.00%"
    if any(token in normalized for token in count_tokens):
        return "0"
    if re.search(r"(?:^|\s)\d+(?:\.\d+)?(?:[-—]\d+(?:\.\d+)?|及以上)(?:$|\s)", normalized):
        return "0"
    return "0.##"


def _style_sheet(sheet, preserve_template_style: bool = False) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = sheet.freeze_panes or "A2"
    used_max_row = max(sheet.max_row, 1)
    used_max_col = max(sheet.max_column, 1)
    merged_title_rows = {
        item.min_row for item in sheet.merged_cells.ranges
        if item.min_row == item.max_row and item.min_col == 1 and item.max_col > 1
    }
    repeated_subheader_rows = {row + 1 for row in merged_title_rows}
    header_by_column: dict[int, str] = {}
    max_lengths = [0] * (used_max_col + 1)
    body_font = Font(color="24384F")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    no_fill = PatternFill(fill_type=None)
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=used_max_row, min_col=1, max_col=used_max_col), start=1
    ):
        populated = [cell for cell in row if cell.value is not None]
        if not populated:
            continue
        first_value = str(row[0].value or "")
        row_values = [str(cell.value or "").strip() for cell in populated]
        semantic_label = " ".join(str(cell.value or "") for cell in row[:2])
        merged_title = row_number in merged_title_rows
        header_tokens = ("方向", "学校", "班级", "班级", "姓名", "学号", "考试", "维度", "学科", "分数段", "参考人数", "人数", "占比")
        repeated_subheader = row_number in repeated_subheader_rows and len(populated) >= 2
        is_header = (
            first_value in {"维度", "考试"}
            or repeated_subheader
            or (len(populated) >= 2 and sum(any(token == value or token in value for token in header_tokens) for value in row_values) >= 2)
        )
        for cell in populated:
            if isinstance(cell.value, str) and (
                cell.value.strip() == "图" or "图方向与常规相反" in cell.value
            ):
                cell.value = None
                continue
            if row_number <= 200:
                max_lengths[cell.column] = max(max_lengths[cell.column], len(str(cell.value)))
            if isinstance(cell.value, (int, float)):
                cell.number_format = _semantic_number_format(
                    f"{semantic_label} {header_by_column.get(cell.column, '')}", semantic_label
                )
            if not preserve_template_style:
                existing_color = cell.fill.fgColor.rgb if cell.fill.fill_type == "solid" else None
                if not existing_color or existing_color[-6:] not in SCORE_LEVEL_FILLS.values():
                    cell.fill = no_fill
                    cell.font = body_font
                cell.alignment = center
                cell.border = THIN_BORDER
        if row_number == 1 or merged_title:
            for cell in populated:
                cell.fill = PatternFill("solid", fgColor=NAVY)
                cell.font = Font(color=WHITE, bold=True)
            sheet.row_dimensions[row_number].height = 28
        elif first_value.startswith(("【", "按分数线", "固定值", "考试基础", "成绩分析", "上线结构", "目标完成", "临界分析", "历次对比", "物理方向", "历史方向")):
            for cell in populated:
                cell.fill = PatternFill("solid", fgColor=PALE_TEAL)
                cell.font = Font(color=NAVY, bold=True)
        elif first_value in {"特控线", "211线", "985线"}:
            for cell in populated:
                cell.fill = PatternFill("solid", fgColor=PALE_GOLD)
                cell.font = Font(color=NAVY, bold=True)
        elif is_header:
            for cell in populated:
                cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
                cell.font = Font(color=NAVY, bold=True)
        for cell in populated:
            if isinstance(cell.value, str) and cell.value.strip():
                header_by_column[cell.column] = cell.value.strip()
    if not preserve_template_style:
        for column in range(1, used_max_col + 1):
            letter = get_column_letter(column)
            sheet.column_dimensions[letter].width = min(max(max_lengths[column] * 1.6 + 3, 10), 26)
    has_merged_title = any(item.min_row == 1 for item in sheet.merged_cells.ranges)
    sheet.auto_filter.ref = None if has_merged_title else (sheet.dimensions if used_max_row > 1 else None)


def _style_workbook(workbook) -> None:
    for sheet in workbook.worksheets:
        _style_sheet(sheet, preserve_template_style=sheet.title in {"物理", "历史"})


def _safe_filename(value: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\r\n]+', "_", str(value)).strip(" ._")
    return cleaned or "未命名"


def _exam(connection, exam_key: str) -> dict:
    row = connection.execute(
        """
        SELECT e.id, e.exam_key, e.name, e.exam_date, e.grade_id, g.name AS grade
        FROM exams e JOIN grades g ON g.id=e.grade_id
        WHERE e.exam_key=? AND e.archived_at IS NULL
        """,
        (exam_key,),
    ).fetchone()
    if row is None:
        raise ValueError("考试不存在")
    return dict(row)


def _selected_exams(connection, exam_key: str, exam_keys: list[str] | None) -> list[dict]:
    """Resolve an explicit export selection in stable date order."""
    selected_keys = list(dict.fromkeys(key for key in (exam_keys or [exam_key]) if key))
    if not selected_keys:
        raise ValueError("至少选择一场考试。")
    placeholders = ",".join("?" for _ in selected_keys)
    rows = connection.execute(
        f"""
        SELECT e.id, e.exam_key, e.name, e.exam_date, e.grade_id, g.name AS grade
        FROM exams e JOIN grades g ON g.id=e.grade_id
        WHERE e.exam_key IN ({placeholders}) AND e.archived_at IS NULL
        ORDER BY e.exam_date, e.id
        """,
        selected_keys,
    ).fetchall()
    if len(rows) != len(selected_keys):
        raise ValueError("所选考试不存在或已归档。")
    return [dict(row) for row in rows]


def _student_rows(connection, exam_key: str) -> list[dict]:
    exam = _exam(connection, exam_key)
    rows = [
        dict(row) for row in connection.execute(
            """
            SELECT er.id AS result_id, er.student_id, er.raw_total, er.scaled_total, er.effective_total,
                   se.track, se.student_type, se.organization_id, se.class_id,
                   c.class_name, o.name AS school,
                   COALESCE(s.display_student_no, s.student_no) AS student_no,
                   COALESCE(s.identity_key, 'legacy|' || s.student_no) AS identity_key, s.name
            FROM exam_results er
            JOIN student_exam_enrollments se ON se.id=er.enrollment_id
            JOIN classes c ON c.id=se.class_id
            JOIN organizations o ON o.id=se.organization_id
            JOIN students s ON s.id=er.student_id
            WHERE er.exam_id=? AND er.attendance_status='normal' AND er.effective_total IS NOT NULL AND se.student_type='普通生'
            ORDER BY se.track, o.name, c.class_name, er.effective_total DESC, s.name
            """,
            (exam["id"],),
        )
    ]
    scores = connection.execute(
        """
        SELECT ss.* FROM subject_scores ss
        JOIN exam_results er ON er.id=ss.exam_result_id
        WHERE er.exam_id=?
        """,
        (exam["id"],),
    ).fetchall()
    by_result: dict[int, dict] = defaultdict(dict)
    for score in scores:
        by_result[score["exam_result_id"]][score["subject_code"]] = dict(score)
    statistics = {track: grade_statistics(connection, exam_key, track) for track in SUBJECT_ORDER}
    rows = [row for row in rows if row["result_id"] in statistics[row["track"]]["ranks"]["__TOTAL__"]]
    for row in rows:
        grade = statistics[row["track"]]
        result_id = row["result_id"]
        row["total_grade_rank_raw"] = grade["raw_ranks"]["__TOTAL__"].get(result_id)
        row["total_grade_rank_scaled"] = grade["ranks"]["__TOTAL__"].get(result_id)
        row["subjects"] = by_result[result_id]
        for code, subject in row["subjects"].items():
            subject["grade_rank_raw"] = grade["raw_ranks"].get(code, {}).get(result_id)
            subject["grade_rank_scaled"] = grade["ranks"].get(code, {}).get(result_id)
    return rows


def _remove_other_sheets(workbook, keep: list[str]) -> None:
    for sheet in list(workbook.worksheets):
        if sheet.title not in keep:
            workbook.remove(sheet)


def _style_row(sheet, source_row: int, target_row: int, max_column: int) -> None:
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    for column in range(1, max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy.copy(source._style)


def _clear_data_rows(sheet, first_row: int = 2) -> None:
    if sheet.max_row >= first_row:
        sheet.delete_rows(first_row, sheet.max_row - first_row + 1)


def _flat_student_row(item: dict) -> list:
    values = [item["school"], item["class_name"], item["name"]]
    for index, code in enumerate(SUBJECT_ORDER[item["track"]]):
        score = item["subjects"].get(code, {})
        if index < 4:
            values.extend([
                score.get("raw_score"), score.get("grade_rank_raw")
            ])
        else:
            values.extend([
                score.get("raw_score"), score.get("grade_rank_raw"), score.get("grade_level"),
                score.get("scaled_score"), score.get("grade_rank_scaled"),
            ])
    values.extend([
        item.get("raw_total"), item.get("total_grade_rank_raw"),
        item.get("effective_total"), item.get("total_grade_rank_scaled"),
    ])
    return values


def _new_class_workbook() -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title in ("物理", "历史"):
        workbook.create_sheet(title)
    for track in ("物理", "历史"):
        sheet = workbook[track]
        headers = ["学校", "班级", "姓名"]
        for index, code in enumerate(SUBJECT_ORDER[track]):
            name = SUBJECTS[code].name
            headers.extend([name, name + "排名"])
            if index >= 4:
                headers.extend([name + "等级", name + "赋分", name + "赋分排名"])
        headers.extend(["总分", "总分排名", "总分赋分", "总分赋分排名"])
        sheet.append(headers)
        for column in range(1, len(headers) + 1):
            letter = get_column_letter(column)
            sheet.column_dimensions[letter].width = 18 if column <= 3 else 13
            sheet.cell(1, column).fill = PatternFill("solid", fgColor=NAVY)
            sheet.cell(1, column).font = Font(color=WHITE, bold=True)
        sheet.row_dimensions[1].height = 30
    return workbook


def _populate_direction_sheets(workbook, rows: list[dict]) -> None:
    for track in ("物理", "历史"):
        if track not in workbook.sheetnames:
            continue
        sheet = workbook[track]
        template_style_row = 2 if sheet.max_row >= 2 else 1
        _clear_data_rows(sheet, 2)
        selected = sorted(
            (
                row for row in rows
                if row["track"] == track 
            ),
            key=lambda row: (-(row["effective_total"] or -1), row["school"], row["class_name"], row["name"]),
        )
        for row_index, item in enumerate(selected, start=2):
            _style_row(sheet, template_style_row, row_index, len(_flat_student_row(item)))
            for column, value in enumerate(_flat_student_row(item), start=1):
                sheet.cell(row_index, column, value)
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{max(1, sheet.max_row)}"
        sheet.freeze_panes = "D2"


def _apply_class_export_borders(workbook) -> None:
    for track in ("物理", "历史"):
        if track not in workbook.sheetnames:
            continue
        sheet = workbook[track]
        for row in sheet.iter_rows(min_row=1, max_row=max(1, sheet.max_row), min_col=1, max_col=sheet.max_column):
            for cell in row:
                cell.border = BLACK_THIN_BORDER


def _export_class_zip_single(
    connection,
    exam_key: str,
    output_folder: str | Path,
    organization_ids: list[int] | None = None,
    class_ids: list[int] | None = None,
) -> Path:
    exam = _exam(connection, exam_key); rows = _student_rows(connection, exam_key)
    selected_organizations = set(organization_ids or [])
    selected_classes = set(class_ids or [])
    grouped: dict[tuple, list[dict]] = defaultdict(list)
    for row in rows:
        if (
            (not selected_organizations or row["organization_id"] in selected_organizations)
            and (not selected_classes or row["class_id"] in selected_classes)
        ):
            grouped[(row["track"], row["school"], row["class_name"])].append(row)
    output_folder = Path(output_folder); output_folder.mkdir(parents=True, exist_ok=True)
    zip_path = output_folder / _safe_filename(f"{exam['grade']}{exam['name']}班级成绩.zip")
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for (track, school, class_name), members in grouped.items():
            workbook = _new_class_workbook()
            _remove_other_sheets(workbook, [track])
            _populate_direction_sheets(workbook, sorted(members, key=lambda row: -(row["effective_total"] or -1)))
            _style_workbook(workbook)
            _apply_class_export_borders(workbook)
            filename = _safe_filename(f"{class_name}班{school}{exam['name']}考试成绩.xlsx")
            temp_path = output_folder / f".{filename}"
            workbook.save(temp_path)
            archive.write(temp_path, filename)
            temp_path.unlink()
    return zip_path


def export_class_zip(
    connection,
    exam_key: str,
    output_folder: str | Path,
    exam_keys: list[str] | None = None,
    organization_ids: list[int] | None = None,
    class_ids: list[int] | None = None,
) -> Path:
    exams = _selected_exams(connection, exam_key, exam_keys)
    generated = [
        (
            exam,
            _export_class_zip_single(
                connection,
                exam["exam_key"],
                output_folder,
                organization_ids,
                class_ids,
            ),
        )
        for exam in exams
    ]
    if len(generated) == 1:
        return generated[0][1]
    output_folder = Path(output_folder)
    zip_path = output_folder / "多考试班级成绩.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for exam, source_path in generated:
            folder = _safe_filename(f"{exam['grade']}{exam['name']}")
            with zipfile.ZipFile(source_path) as source:
                for name in source.namelist():
                    archive.writestr(f"{folder}/{name}", source.read(name))
    return zip_path


