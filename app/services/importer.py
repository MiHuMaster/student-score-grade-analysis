from __future__ import annotations

import hashlib
import json
import logging
import math
import re
import sqlite3
import unicodedata
import uuid
from collections import defaultdict
from copy import copy
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path
from time import perf_counter
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import from_excel

from app.config import (
    CORE_RULE_VERSION,
    SCORE_SHEETS,
    SUBJECTS,
    SUBJECT_NAME_TO_CODE,
    TRACKS,
)


LOGGER = logging.getLogger(__name__)


@dataclass
class ImportIssue:
    severity: str
    issue_code: str
    message: str
    suggestion: str | None = None
    sheet_name: str | None = None
    row_number: int | None = None
    field_name: str | None = None
    raw_value: str | None = None


@dataclass
class WorkbookPayload:
    schools: list[dict] = field(default_factory=list)
    students: list[dict] = field(default_factory=list)
    scores: list[dict] = field(default_factory=list)
    score_lines: list[dict] = field(default_factory=list)
    issues: list[ImportIssue] = field(default_factory=list)
    sheets_found: list[str] = field(default_factory=list)
    header_mappings: list[dict] = field(default_factory=list)
    timings: dict[str, float] = field(default_factory=dict)


def build_exam_key(grade: str, exam_name: str, exam_date: date) -> str:
    normalized_name = re.sub(r"\s+", "", exam_name).strip().lower()
    raw = f"{grade.strip()}|{normalized_name}|{exam_date.isoformat()}"
    digest = hashlib.sha256(raw.encode("utf-8")).hexdigest()[:16]
    return f"exam_{digest}"


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _normalized_text(value: Any) -> str:
    """Normalize spreadsheet identity text without changing meaningful characters."""
    text = unicodedata.normalize("NFKC", _text(value))
    return re.sub(r"\s+", "", text)


def _student_number(value: Any) -> str:
    """Keep leading-zero text IDs while removing Excel's accidental numeric '.0'."""
    if isinstance(value, bool) or value is None:
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value))
    text = unicodedata.normalize("NFKC", _text(value))
    if re.fullmatch(r"[+-]?\d+\.0+", text):
        return text.split(".", 1)[0]
    return re.sub(r"\s+", "", text)


def _regular_identity_key(student_no: Any, name: Any) -> str:
    """Return the permanent ordinary-student identity: normalized number + name."""
    return "regular|" + "|".join(
        (
            _student_number(student_no).casefold(),
            _normalized_text(name).casefold(),
        )
    )




IDENTITY_HEADER_ALIASES = {
    "年级": "grade", "年级名称": "grade",
    "考试": "exam_name", "考试名": "exam_name", "考试名称": "exam_name",
    "日期": "exam_date", "考试日期": "exam_date", "考试时间": "exam_date",
    "学生编号": "student_no", "考号": "student_no", "学号": "student_no",
    "班级": "class_name",
    "学生姓名": "student_name", "姓名": "student_name",
    "学生类型": "student_type", "类型": "student_type",
}

IDENTITY_CANONICAL_HEADERS = {
    "grade": "年级", "exam_name": "考试名称", "exam_date": "考试时间",
    "student_no": "学号", "class_name": "班级", "student_name": "姓名",
    "student_type": "学生类型",
}

def _header_token(value: Any) -> str:
    text = unicodedata.normalize("NFKC", _text(value))
    return re.sub(r"\s+", "", text).replace("（", "(").replace("）", ")")


def _canonical_subject_header(code: str, field_name: str) -> str:
    subject = SUBJECTS[code].name
    return {
        "raw_score": subject,
        "network_rank_raw": f"{subject}排名",
        "grade_level": f"{subject}等级",
        "scaled_score": f"{subject}赋分",
        "network_rank_scaled": f"{subject}赋分排名",
    }[field_name]


def _canonical_total_header(field_name: str) -> str:
    return {
        "raw_total": "总分",
        "network_rank_raw": "总分排名",
        "scaled_total": "总分赋分",
        "network_rank_scaled": "总分赋分排名",
    }[field_name]


def _build_score_column_map(sheet, sheet_name: str, payload: WorkbookPayload) -> dict:
    mapping = {
        "identity": {},
        "subjects": {code: {} for code in SUBJECTS},
        "totals": {},
        "max_column": sheet.max_column,
    }
    current_context: tuple[str, str] | None = None

    def assign(target: str, field_name: str, column: int, original: str, canonical: str, code: str | None = None) -> None:
        bucket = mapping[target] if code is None else mapping[target][code]
        if field_name in bucket and bucket[field_name] != column:
            payload.issues.append(
                ImportIssue(
                    "error", "score_header_duplicate",
                    f"表头“{canonical}”对应了多个列，无法唯一确定数据来源",
                    "删除重复表头，或在表头中补充学科与原始分/赋分含义",
                    sheet_name, 1, original, original,
                )
            )
            return
        bucket[field_name] = column
        compact_original = re.sub(r"\s+", "", original)
        if original and unicodedata.normalize("NFKC", compact_original) != unicodedata.normalize("NFKC", canonical):
            payload.header_mappings.append(
                {"sheet": sheet_name, "column": column, "original": original, "normalized": canonical}
            )

    for column in range(1, sheet.max_column + 1):
        original = _text(sheet.cell(1, column).value)
        token = _header_token(original)
        if not token:
            continue
        if token in IDENTITY_HEADER_ALIASES:
            field_name = IDENTITY_HEADER_ALIASES[token]
            assign("identity", field_name, column, original, IDENTITY_CANONICAL_HEADERS[field_name])
            current_context = None
            continue
        matched = False
        for subject_name, code in SUBJECT_NAME_TO_CODE.items():
            if token in {subject_name, f"{subject_name}成绩", f"{subject_name}原始分"}:
                assign("subjects", "raw_score", column, original, _canonical_subject_header(code, "raw_score"), code)
                current_context = (code, "raw")
                matched = True
                break
            patterns = (
                (rf"{re.escape(subject_name)}(?:N|赋分)", "scaled_score", "scaled"),
                (rf"{re.escape(subject_name)}(?:M|等级)", "grade_level", "scaled"),
                (rf"{re.escape(subject_name)}(?:N排名|赋分排名)", "network_rank_scaled", "scaled"),
                (rf"{re.escape(subject_name)}(?:排名|年级排名)", "network_rank_raw", "raw"),
            )
            for pattern, field_name, mode in patterns:
                if re.fullmatch(pattern, token, re.IGNORECASE):
                    assign("subjects", field_name, column, original, _canonical_subject_header(code, field_name), code)
                    current_context = (code, mode)
                    matched = True
                    break
            if matched:
                break
        if matched:
            continue

        if token in {"总分", "总成绩", "原始总分", "总分原始分"}:
            assign("totals", "raw_total", column, original, _canonical_total_header("raw_total"))
            current_context = ("total", "raw")
        elif re.fullmatch(r"总分(?:N|赋分)|赋分总分", token, re.IGNORECASE):
            assign("totals", "scaled_total", column, original, _canonical_total_header("scaled_total"))
            current_context = ("total", "scaled")
        elif token in {"总分排名", "总分年级排名"}:
            assign("totals", "network_rank_raw", column, original, _canonical_total_header("network_rank_raw"))
            current_context = ("total", "raw")
        elif re.fullmatch(r"总分(?:N排名|赋分排名)", token, re.IGNORECASE):
            assign("totals", "network_rank_scaled", column, original, _canonical_total_header("network_rank_scaled"))
            current_context = ("total", "scaled")
        elif token == "排名" and current_context:
            owner, mode = current_context
            field_name = "network_rank_scaled" if mode == "scaled" else "network_rank_raw"
            if owner == "total":
                assign("totals", field_name, column, original, _canonical_total_header(field_name))
            else:
                assign("subjects", field_name, column, original, _canonical_subject_header(owner, field_name), owner)
    return mapping


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if math.isnan(float(value)) or math.isinf(float(value)):
            return None
        return float(value)
    text = str(value).strip().replace(",", "")
    if text in {"", "/", "#N/A", "#VALUE!", "#REF!", "（手动输入）", "(手动输入)"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _integer(value: Any) -> int | None:
    number = _number(value)
    return None if number is None else int(round(number))


def _parse_date(value: Any, epoch: datetime) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 20000 <= float(value) <= 80000:
        converted = from_excel(value, epoch)
        return converted.date() if isinstance(converted, datetime) else converted
    text = _text(value)
    if not text:
        return None
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    return None


PREPARED_PREFIX_HEADERS = ["年级", "考试名称", "考试时间", "学校", "学号"]
REFERENCE_SHEETS = ["学校信息", "学生信息"]
TOTAL_SCORE_SHEETS = ["物理方向学生成绩", "历史方向学生成绩"]
OPTIONAL_SCORE_SHEETS = ["分数线表"]

SCHOOL_SHEET_ALIASES = ("学校信息", "学校基础信息表")
STUDENT_SHEET_ALIASES = ("学生信息", "学生基础信息")
BASE_SCHOOL_HEADERS = {
    "省份": ("省份", "学校省份"), "学校": ("学校", "学校名称"),
    "年级": ("年级",), "班级": ("班级",),
    "方向": ("方向",), "班级人数": ("班级人数",),
    "特控目标": ("特控目标", "班级特控目标人数", "特控目标人数"),
}
BASE_STUDENT_HEADERS = {
    "姓名": ("姓名", "学生姓名"), "学号": ("学号", "学生编号", "考号"),
    "年级": ("年级",), "班级": ("班级",),
    "学校": ("学校", "学校名称"), "选考标签": ("选考标签", "选科标签"),
}


def _find_sheet(workbook, aliases: tuple[str, ...]):
    return next((workbook[name] for name in aliases if name in workbook.sheetnames), None)


def _aliased_headers(sheet, aliases: dict[str, tuple[str, ...]]) -> dict[str, int]:
    values = {_header_token(cell.value): index for index, cell in enumerate(sheet[1], start=1)}
    return {
        canonical: next((values[name] for name in names if name in values), None)
        for canonical, names in aliases.items()
    }


def _target_columns(sheet, headers: dict[str, int]) -> list[tuple[str, int, int]]:
    """Return dynamic ``xx目标`` columns in their original spreadsheet order."""
    columns: list[tuple[str, int]] = []
    legacy_control = headers.get("特控目标")
    if legacy_control:
        columns.append(("特控", legacy_control))
    for column, cell in enumerate(sheet[1], start=1):
        if column < 7:
            continue
        token = _header_token(cell.value)
        if not token.endswith("目标") or token == "目标":
            continue
        columns.append((token[:-2], column))
    seen: set[str] = set()
    ordered = []
    for name, column in sorted(columns, key=lambda item: item[1]):
        if name in seen:
            continue
        seen.add(name)
        ordered.append((name, column, len(ordered)))
    return ordered


def _school_targets(school: dict) -> list[dict]:
    targets = list(school.get("targets") or [])
    if not targets and school.get("control_target") is not None:
        targets = [{"name": "特控", "order": 0, "count": school["control_target"]}]
    return targets


def _validate_import_context(grade: str, exam_name: str, exam_month: str) -> date:
    if not re.fullmatch(r"高20\d{2}届", grade):
        raise ValueError("年级格式应为“高202x届”")
    if not exam_name or len(exam_name) > 80:
        raise ValueError("请填写 1—80 个字符的考试名称")
    matched = re.fullmatch(r"(20\d{2})年(0?[1-9]|1[0-2])月", exam_month)
    if matched is None:
        raise ValueError("考试时间格式应为“20xx年xx月”")
    return date(int(matched.group(1)), int(matched.group(2)), 1)


def _copy_sheet_layout(source, target) -> None:
    for key, dimension in source.column_dimensions.items():
        target.column_dimensions[key].width = dimension.width
        target.column_dimensions[key].hidden = dimension.hidden
    for key, dimension in source.row_dimensions.items():
        target.row_dimensions[key].height = dimension.height
        target.row_dimensions[key].hidden = dimension.hidden
    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.sheet_properties = copy(source.sheet_properties)
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)
    target.sheet_format = copy(source.sheet_format)
    if source.auto_filter.ref:
        target.auto_filter.ref = source.auto_filter.ref


def _copy_cell_style(source_cell, target_cell) -> None:
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)
    target_cell.number_format = source_cell.number_format
    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.protection = copy(source_cell.protection)


def _snapshot_value(formula_cell, value_cell) -> Any:
    value = formula_cell.value
    if isinstance(value, str) and value.startswith("="):
        return value_cell.value
    return value


def _copy_sheet_snapshot(formula_sheet, value_sheet, target) -> None:
    _copy_sheet_layout(formula_sheet, target)
    for row in formula_sheet.iter_rows():
        for source_cell in row:
            target_cell = target.cell(source_cell.row, source_cell.column)
            target_cell.value = _snapshot_value(
                source_cell, value_sheet.cell(source_cell.row, source_cell.column)
            )
            _copy_cell_style(source_cell, target_cell)
    for merged_range in formula_sheet.merged_cells.ranges:
        target.merge_cells(str(merged_range))


def _copy_reference_sheet_compact(formula_sheet, value_sheet, target) -> None:
    """Copy saved base data without reproducing thousands of non-business cell styles."""
    _copy_sheet_layout(formula_sheet, target)
    for column, source_cell in enumerate(formula_sheet[1], start=1):
        target_cell = target.cell(1, column, source_cell.value)
        _copy_cell_style(source_cell, target_cell)
    for row in value_sheet.iter_rows(min_row=2, values_only=True):
        target.append(list(row))


def _reference_student_indexes(students: list[dict]) -> tuple[list[dict], dict[str, list[dict]], dict[tuple[str, str], list[dict]]]:
    by_number: dict[str, list[dict]] = {}
    by_name_class: dict[tuple[str, str], list[dict]] = {}
    for student in students:
        if student["student_no"]:
            by_number.setdefault(student["student_no"], []).append(student)
        key = (
            _normalized_text(student["name"]).casefold(),
            _normalized_text(student["class_name"]).casefold(),
        )
        if all(key):
            by_name_class.setdefault(key, []).append(student)
    return students, by_number, by_name_class


def _reference_students(reference_workbook) -> tuple[list[dict], dict[str, list[dict]], dict[tuple[str, str], list[dict]]]:
    sheet = _find_sheet(reference_workbook, STUDENT_SHEET_ALIASES)
    if sheet is None:
        raise ValueError("基础信息表缺少“学生信息”子表")
    headers = _aliased_headers(sheet, BASE_STUDENT_HEADERS)
    required = ("姓名", "学号", "班级", "学校")
    missing = [header for header in required if not headers[header]]
    if missing:
        raise ValueError(f"学生信息缺少必填列：{'、'.join(missing)}")
    students: list[dict] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(_text(value) for value in row):
            continue
        student = {
            "row_number": row_number,
            "name": _text(row[headers["姓名"] - 1]),
            "student_no": _student_number(row[headers["学号"] - 1]),
            "class_name": _text(row[headers["班级"] - 1]),
            "school": _text(row[headers["学校"] - 1]),
        }
        students.append(student)
    return _reference_student_indexes(students)


def saved_base_reference_data(connection: sqlite3.Connection) -> dict[str, list[dict]]:
    schools = [dict(row) for row in connection.execute(
        "SELECT id, province, school_name, grade_name, class_name, track, planned_size, control_target FROM base_school_classes ORDER BY school_name, grade_name, class_name"
    ).fetchall()]
    students = [
        {
            "name": row["student_name"], "student_no": row["student_no"],
            "class_name": row["class_name"], "school": row["school_name"],
            "grade": row["grade_name"], "subject_tag": row["subject_tag"] or "",
        }
        for row in connection.execute(
            "SELECT school_name, grade_name, class_name, student_name, student_no, subject_tag FROM base_students ORDER BY school_name, grade_name, class_name, student_no"
        ).fetchall()
    ]
    target_rows = connection.execute(
        """SELECT bct.base_class_id, bct.target_name, bct.target_order, bct.target_count
           FROM base_class_targets bct
           ORDER BY bct.target_order, bct.target_name"""
    ).fetchall()
    targets_by_class: dict[int, dict[str, int]] = defaultdict(dict)
    target_orders: dict[str, int] = {}
    for row in target_rows:
        targets_by_class[row["base_class_id"]][row["target_name"]] = row["target_count"]
        target_orders[row["target_name"]] = min(
            target_orders.get(row["target_name"], row["target_order"]), row["target_order"]
        )
    for row in schools:
        row["targets"] = targets_by_class.get(row["id"], {})
    if not students:
        raise ValueError("尚未保存学生基础信息，请先上传基础信息表")
    target_names = sorted(target_orders, key=lambda name: (target_orders[name], name))
    return {"schools": schools, "students": students, "target_names": target_names}


def _write_saved_base_reference(output, reference_data: dict[str, list[dict]]) -> None:
    school_sheet = output.create_sheet("学校信息")
    target_names = reference_data.get("target_names") or ["特控"]
    school_sheet.append(["省份", "学校", "年级", "班级", "方向", "班级人数", *[f"{name}目标" for name in target_names]])
    for row in reference_data["schools"]:
        targets = row.get("targets") or {"特控": row.get("control_target")}
        school_sheet.append([row["province"], row["school_name"], row["grade_name"], row["class_name"], row.get("track"), row["planned_size"], *[targets.get(name) for name in target_names]])
    student_sheet = output.create_sheet("学生信息")
    student_sheet.append(["学校", "年级", "班级", "学生姓名", "学号", "选科标签"])
    for row in reference_data["students"]:
        student_sheet.append([row["school"], row["grade"], row["class_name"], row["name"], row["student_no"], row["subject_tag"]])


def _prepare_direction_sheet(
    formula_sheet,
    value_sheet,
    target,
    grade: str,
    exam_name: str,
    exam_date: date,
    by_number: dict[str, list[dict]],
    by_name_class: dict[tuple[str, str], list[dict]],
) -> dict:
    started = perf_counter()
    raw_headers = [_header_token(cell.value) for cell in formula_sheet[1]]
    class_columns = [index for index, header in enumerate(raw_headers, start=1) if header in {"班级"}]
    name_columns = [index for index, header in enumerate(raw_headers, start=1) if header in {"姓名", "学生姓名"}]
    number_columns = [index for index, header in enumerate(raw_headers, start=1) if header in {"学号", "学生编号", "考号"}]
    school_columns = [index for index, header in enumerate(raw_headers, start=1) if header == "学校"]
    if len(class_columns) != 1 or len(name_columns) != 1:
        raise ValueError(f"{formula_sheet.title}必须且只能有一个班级列和姓名列")
    map_payload = WorkbookPayload()
    column_map = _build_score_column_map(formula_sheet, formula_sheet.title, map_payload)
    primary = "physics" if "物理" in formula_sheet.title else "history"
    visible_columns = [("班级", class_columns[0]), ("姓名", name_columns[0])]
    for code in ("chinese", "math", "english", primary):
        for field_name in ("raw_score", "network_rank_raw"):
            visible_columns.append((_canonical_subject_header(code, field_name), column_map["subjects"][code].get(field_name)))
    for code in ("chemistry", "biology", "politics", "geography"):
        for field_name in ("raw_score", "network_rank_raw", "grade_level", "scaled_score", "network_rank_scaled"):
            visible_columns.append((_canonical_subject_header(code, field_name), column_map["subjects"][code].get(field_name)))
    for field_name in ("raw_total", "network_rank_raw", "scaled_total", "network_rank_scaled"):
        visible_columns.append((_canonical_total_header(field_name), column_map["totals"].get(field_name)))
    _copy_sheet_layout(formula_sheet, target)
    target.freeze_panes = "C2"
    style_fallback = class_columns[0]
    for target_column, (header, source_column) in enumerate(visible_columns, start=1):
        source_cell = formula_sheet.cell(1, source_column or style_fallback)
        target_cell = target.cell(1, target_column, header)
        _copy_cell_style(source_cell, target_cell)
        target.column_dimensions[target_cell.column_letter].width = formula_sheet.column_dimensions[source_cell.column_letter].width or 15
    metadata_start = len(visible_columns) + 1
    for offset, header in enumerate(PREPARED_PREFIX_HEADERS, start=metadata_start):
        target_cell = target.cell(1, offset, header)
        _copy_cell_style(formula_sheet.cell(1, style_fallback), target_cell)
        target.column_dimensions[target_cell.column_letter].hidden = True
    # 数据行共享列样式；样式对象在 openpyxl 中不可变，复用不会污染原工作簿，
    # 也避免按“行 × 列”反复复制字体、边框和填充。
    data_styles = []
    for target_column, (_header, source_column) in enumerate(visible_columns, start=1):
        style_cell = target.cell(2, target_column)
        _copy_cell_style(formula_sheet.cell(2, source_column or style_fallback), style_cell)
        data_styles.append(copy(style_cell._style))
    metadata_style_cell = target.cell(2, metadata_start)
    _copy_cell_style(formula_sheet.cell(2, style_fallback), metadata_style_cell)
    metadata_style = copy(metadata_style_cell._style)

    matched = 0
    unmatched_rows: list[int] = []
    ambiguous_rows: list[int] = []
    formula_rows = formula_sheet.iter_rows(min_row=2, max_col=formula_sheet.max_column)
    value_rows = value_sheet.iter_rows(min_row=2, max_col=value_sheet.max_column)
    for row_number, (formula_cells, value_cells) in enumerate(zip(formula_rows, value_rows), start=2):
        formula_values = [cell.value for cell in formula_cells]
        value_values = [cell.value for cell in value_cells]
        if not any(_text(value) for value in value_values) and not any(_text(value) for value in formula_values):
            continue
        student_no = _student_number(value_values[number_columns[0] - 1]) if number_columns else ""
        class_name = _text(value_values[class_columns[0] - 1])
        student_name = _text(value_values[name_columns[0] - 1])
        candidates = by_number.get(student_no, []) if student_no else []
        if candidates and (student_name or class_name):
            exact_number = [
                item for item in candidates
                if (not student_name or _normalized_text(item["name"]).casefold() == _normalized_text(student_name).casefold())
                and (not class_name or _normalized_text(item["class_name"]).casefold() == _normalized_text(class_name).casefold())
            ]
            candidates = exact_number
        if len(candidates) != 1:
            candidates = by_name_class.get(
                (
                    _normalized_text(student_name).casefold(),
                    _normalized_text(class_name).casefold(),
                ),
                [],
            )
        master = candidates[0] if len(candidates) == 1 else None
        if master:
            matched += 1
        elif len(candidates) > 1:
            ambiguous_rows.append(row_number)
        else:
            unmatched_rows.append(row_number)
        metadata = [
            grade,
            exam_name,
            exam_date,
            master["school"] if master else "",
            master["student_no"] if master else "",
        ]
        for target_column, (_header, source_column) in enumerate(visible_columns, start=1):
            target_cell = target.cell(row_number, target_column)
            target_cell.value = _snapshot_value(
                formula_cells[source_column - 1], value_cells[source_column - 1]
            ) if source_column else None
            target_cell._style = data_styles[target_column - 1]
        for target_column, value in enumerate(metadata, start=metadata_start):
            target_cell = target.cell(row_number, target_column, value)
            target_cell._style = metadata_style
            if target_column == metadata_start + 2:
                target_cell.number_format = 'yyyy"年"mm"月"'
    return {
        "sheet": formula_sheet.title,
        "matched_rows": matched,
        "unmatched_rows": len(unmatched_rows),
        "ambiguous_rows": len(ambiguous_rows),
        "unmatched_samples": unmatched_rows[:20],
        "ambiguous_samples": ambiguous_rows[:20],
        "elapsed_seconds": round(perf_counter() - started, 3),
    }






def prepare_import_workbook(
    reference_path: str | Path | None,
    score_path: str | Path,
    output_path: str | Path,
    grade: str,
    exam_name: str,
    exam_month: str,
    reference_data: dict[str, list[dict]] | None = None,
) -> dict:
    """Create a value-stable, matched workbook for the existing validation pipeline."""
    started = perf_counter()
    grade = _text(grade)
    exam_name = _text(exam_name)
    exam_month = _text(exam_month)
    exam_date = _validate_import_context(grade, exam_name, exam_month)
    reference_formula = load_workbook(reference_path, data_only=False) if reference_path else None
    reference_values = load_workbook(reference_path, data_only=True) if reference_path else None
    score_formula = load_workbook(score_path, data_only=False)
    score_values = load_workbook(score_path, data_only=True)
    workbook_load_seconds = perf_counter() - started
    try:
        missing_reference = []
        if reference_data is None:
            if _find_sheet(reference_formula, SCHOOL_SHEET_ALIASES) is None:
                missing_reference.append("学校信息")
            if _find_sheet(reference_formula, STUDENT_SHEET_ALIASES) is None:
                missing_reference.append("学生信息")
        missing_scores = [name for name in TOTAL_SCORE_SHEETS if name not in score_formula.sheetnames]
        if missing_reference:
            raise ValueError(f"基础信息对照表缺少工作表：{'、'.join(missing_reference)}")
        if missing_scores:
            raise ValueError(f"学生总体赋分成绩表缺少工作表：{'、'.join(missing_scores)}")
        _, by_number, by_name_class = (
            _reference_student_indexes(reference_data["students"])
            if reference_data is not None else _reference_students(reference_values)
        )
        output = Workbook()
        output.remove(output.active)
        base_copy_started = perf_counter()
        if reference_data is not None:
            _write_saved_base_reference(output, reference_data)
        else:
            for output_name, aliases in (("学校信息", SCHOOL_SHEET_ALIASES), ("学生信息", STUDENT_SHEET_ALIASES)):
                target = output.create_sheet(output_name)
                _copy_reference_sheet_compact(
                    _find_sheet(reference_formula, aliases), _find_sheet(reference_values, aliases), target
                )
        base_copy_seconds = perf_counter() - base_copy_started
        matching: list[dict] = []
        for sheet_name in [
            *TOTAL_SCORE_SHEETS,
            *(name for name in OPTIONAL_SCORE_SHEETS if name in score_formula.sheetnames),
        ]:
            target = output.create_sheet(sheet_name)
            if sheet_name in SCORE_SHEETS:
                matching.append(
                    _prepare_direction_sheet(
                        score_formula[sheet_name], score_values[sheet_name], target,
                        grade, exam_name, exam_date, by_number, by_name_class,
                    )
                )
            else:
                _copy_sheet_snapshot(score_formula[sheet_name], score_values[sheet_name], target)
        output.calculation.fullCalcOnLoad = False
        output.calculation.forceFullCalc = False
        output.calculation.calcMode = "manual"
        save_started = perf_counter()
        output.save(output_path)
        workbook_save_seconds = perf_counter() - save_started
    finally:
        if reference_formula is not None:
            reference_formula.close()
        if reference_values is not None:
            reference_values.close()
        score_formula.close()
        score_values.close()
    return {
        "grade": grade,
        "exam_name": exam_name,
        "exam_month": exam_month,
        "exam_date": exam_date.isoformat(),
        "matching": matching,
        "matched_rows": sum(item["matched_rows"] for item in matching),
        "unmatched_rows": sum(item["unmatched_rows"] for item in matching),
        "ambiguous_rows": sum(item["ambiguous_rows"] for item in matching),
        "timings": {
            "工作簿读取": round(workbook_load_seconds, 3),
            "基础信息写入": round(base_copy_seconds, 3),
            "匹配与规范表写入": round(sum(item.get("elapsed_seconds", 0) for item in matching), 3),
            "规范表保存": round(workbook_save_seconds, 3),
            "规范表生成合计": round(perf_counter() - started, 3),
        },
    }


def _subject_codes_from_tag(tag: str, track: str) -> list[str]:
    character_map = {
        "物": "physics",
        "历": "history",
        "史": "history",
        "化": "chemistry",
        "生": "biology",
        "政": "politics",
        "地": "geography",
    }
    primary = "physics" if track == "物理" else "history"
    codes = []
    for character in tag:
        code = character_map.get(character)
        if code and code != primary and code not in codes:
            codes.append(code)
    return codes


def _parse_school_sheet(workbook, payload: WorkbookPayload) -> None:
    sheet = _find_sheet(workbook, SCHOOL_SHEET_ALIASES)
    if sheet is None:
        payload.issues.append(
            ImportIssue("warning", "missing_school_sheet", "缺少学校基础信息表，班级目标将无法导入")
        )
        return
    headers = _aliased_headers(sheet, BASE_SCHOOL_HEADERS)
    target_columns = _target_columns(sheet, headers)
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        school_name = _text(row[headers.get("学校", 0) - 1]) if headers.get("学校") else ""
        if not school_name:
            continue
        track_text = _text(row[headers.get("方向", 0) - 1]) if headers.get("方向") else ""
        class_name = _text(row[headers.get("班级", 0) - 1]) if headers.get("班级") else ""
        targets = []
        for target_name, column, target_order in target_columns:
            raw_value = row[column - 1] if column <= len(row) else None
            if not _text(raw_value):
                continue
            numeric = _number(raw_value)
            if numeric is None or numeric < 0 or not math.isclose(numeric, round(numeric)):
                payload.issues.append(
                    ImportIssue(
                        "error", "target_count_invalid",
                        f"{target_name}目标必须是非负整数",
                        "请填写整数人数或留空", sheet.title, row_number,
                        f"{target_name}目标", _text(raw_value),
                    )
                )
                continue
            targets.append({"name": target_name, "order": target_order, "count": int(round(numeric))})
        control_target = next((item["count"] for item in targets if item["name"] == "特控"), None)
        payload.schools.append(
            {
                "province": _text(row[headers.get("省份", 0) - 1]) if headers.get("省份") else "",
                "school": school_name,
                "grade": _text(row[headers.get("年级", 0) - 1]) if headers.get("年级") else "",
                "class_name": class_name,
                "track": track_text if track_text in {"物理", "历史"} else None,
                "planned_size": _integer(row[headers.get("班级人数", 0) - 1]) if headers.get("班级人数") else None,
                "control_target": control_target,
                "targets": targets,
                "row_number": row_number,
            }
        )


def _parse_student_sheet(workbook, payload: WorkbookPayload) -> None:
    sheet = _find_sheet(workbook, STUDENT_SHEET_ALIASES)
    if sheet is None:
        payload.issues.append(
            ImportIssue("error", "missing_student_sheet", "缺少学生信息工作表", "请使用标准模板并补齐学生信息")
        )
        return
    headers = _aliased_headers(sheet, BASE_STUDENT_HEADERS)
    required = ("姓名", "学号", "年级", "班级", "学校")
    missing = [header for header in required if not headers[header]]
    if missing:
        payload.issues.append(
            ImportIssue(
                "error",
                "student_headers_missing",
                f"学生信息缺少必填列：{'、'.join(missing)}",
                "请恢复标准表头",
                sheet.title,
                1,
            )
        )
        return
    seen: set[str] = set()
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        name = _text(row[headers["姓名"] - 1])
        student_no = _student_number(row[headers["学号"] - 1])
        if not any(_text(value) for value in row):
            continue
        if not name:
            payload.issues.append(
                ImportIssue("error", "student_name_missing", "学生姓名为空", "补充姓名", sheet.title, row_number, "姓名")
            )
        if not student_no:
            payload.issues.append(
                ImportIssue("error", "student_no_missing", "学号为空，无法建立稳定学生身份", "补充学校提供的唯一学号", sheet.title, row_number, "学号")
            )
        school_name = _text(row[headers["学校"] - 1])
        grade_name = _text(row[headers["年级"] - 1])
        class_name = _text(row[headers["班级"] - 1])
        identity_key = _regular_identity_key(student_no, name)
        if student_no and identity_key in seen:
            payload.issues.append(
                ImportIssue(
                    "error", "student_identity_duplicate",
                    "学生信息中学号和姓名组合重复",
                    "核对重复学生记录；学校、班级和方向不参与永久身份", sheet.title, row_number, "学号", student_no,
                )
            )
        if student_no:
            seen.add(identity_key)
        payload.students.append(
            {
                "student_no": student_no,
                "name": name,
                "grade": grade_name,
                "class_name": class_name,
                "school": school_name,
                "subject_tag": _text(row[headers["选考标签"] - 1]) if headers["选考标签"] else "",
                "student_type": "普通生",
                "identity_key": identity_key,
                "row_number": row_number,
            }
        )


def parse_base_information_workbook(path: str | Path) -> WorkbookPayload:
    """Parse the stable two-sheet base-information workbook without requiring scores."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        payload = WorkbookPayload(sheets_found=list(workbook.sheetnames))
        _parse_school_sheet(workbook, payload)
        _parse_student_sheet(workbook, payload)
        _validate_single_school(payload)
        return payload
    finally:
        workbook.close()


HISTORICAL_CLASS_TRACK_MESSAGE = (
    "该班级已存在历史考试数据，班级方向不能直接修改，历史考试不会被自动覆盖。"
)


def _class_has_formal_references(connection: sqlite3.Connection, class_id: int) -> bool:
    return connection.execute(
        """SELECT EXISTS(
             SELECT 1 FROM student_exam_enrollments WHERE class_id=?
           ) OR EXISTS(
             SELECT 1 FROM class_targets WHERE class_id=?
           ) OR EXISTS(
             SELECT 1 FROM class_stage_targets WHERE class_id=?
           ) AS has_references""",
        (class_id, class_id, class_id),
    ).fetchone()["has_references"] == 1


def save_base_information(connection: sqlite3.Connection, payload: WorkbookPayload) -> dict:
    _validate_single_school(payload, connection)
    errors = [issue for issue in payload.issues if issue.severity == "error"]
    if errors:
        raise ValueError("基础信息存在错误，请先修正后再保存")
    # 基础信息文件是年级级快照。旧实现只做 UPSERT，会让已经从新文件中删除或
    # 改号的学生永久残留，随后在无学号成绩表中造成“姓名+班级”伪重复。
    # 只替换本次文件实际涉及的年级，其他年级完全保留。
    grades = sorted({
        _text(item.get("grade"))
        for item in [*payload.schools, *payload.students]
        if _text(item.get("grade"))
    })
    try:
        connection.execute("BEGIN IMMEDIATE")
        for grade in grades:
            connection.execute("DELETE FROM base_students WHERE grade_name=?", (grade,))
            connection.execute("DELETE FROM base_school_classes WHERE grade_name=?", (grade,))
            connection.execute(
                """DELETE FROM class_stage_targets
                   WHERE class_id IN (
                     SELECT c.id FROM classes c JOIN grades g ON g.id=c.grade_id
                     WHERE g.name=?
                   )""",
                (grade,),
            )
        for school in payload.schools:
            if not all((school.get("school"), school.get("grade"), school.get("class_name"))):
                continue
            connection.execute(
                """INSERT INTO base_school_classes(
                     province, school_name, grade_name, class_name, track, planned_size, control_target
                   ) VALUES (?, ?, ?, ?, ?, ?, ?)
                   ON CONFLICT(school_name, grade_name, class_name) DO UPDATE SET
                     province=excluded.province, track=excluded.track,
                     planned_size=excluded.planned_size, control_target=excluded.control_target,
                     updated_at=CURRENT_TIMESTAMP""",
                (school.get("province") or None, school["school"], school["grade"], school["class_name"],
                 school.get("track"), school.get("planned_size"), school.get("control_target")),
            )
            base_class_id = connection.execute(
                """SELECT id FROM base_school_classes
                   WHERE school_name=? AND grade_name=? AND class_name=?""",
                (school["school"], school["grade"], school["class_name"]),
            ).fetchone()["id"]
            for target in _school_targets(school):
                connection.execute(
                    """INSERT INTO base_class_targets(
                         base_class_id, target_name, target_order, target_count
                       ) VALUES (?, ?, ?, ?)""",
                    (base_class_id, target["name"], target["order"], target["count"]),
                )
            if school.get("track") in TRACKS:
                organization_id = _get_or_create_id(
                    connection, "organizations", {"name": school["school"]},
                    {"province": school.get("province") or None},
                )
                grade_id = _get_or_create_id(connection, "grades", {"name": school["grade"]})
                existing_classes = connection.execute(
                    """SELECT id, track FROM classes
                       WHERE organization_id=? AND grade_id=? AND class_name=?
                         AND archived_at IS NULL""",
                    (organization_id, grade_id, school["class_name"]),
                ).fetchall()
                opposite_ids = [
                    row["id"] for row in existing_classes if row["track"] != school["track"]
                ]
                if any(_class_has_formal_references(connection, class_id) for class_id in opposite_ids):
                    raise ValueError(HISTORICAL_CLASS_TRACK_MESSAGE)
                class_id = _get_or_create_id(
                    connection, "classes",
                    {
                        "organization_id": organization_id,
                        "grade_id": grade_id,
                        "class_name": school["class_name"],
                        "track": school["track"],
                    },
                )
                connection.execute(
                    """UPDATE classes SET planned_size=?, control_target=?, archived_at=NULL
                       WHERE id=?""",
                    (school.get("planned_size"), school.get("control_target"), class_id),
                )
                for target in _school_targets(school):
                    connection.execute(
                        """INSERT INTO class_stage_targets(
                             class_id, target_name, target_order, target_count, source
                           ) VALUES (?, ?, ?, ?, 'base')
                           ON CONFLICT(class_id, target_name) DO UPDATE SET
                             target_order=excluded.target_order,
                             target_count=excluded.target_count,
                             source='base', updated_at=CURRENT_TIMESTAMP""",
                        (class_id, target["name"], target["order"], target["count"]),
                    )
                # Direction is immutable once a class has formal history. Only an unused
                # opposite-direction class may be archived; historical enrollments are never moved.
                for opposite_id in opposite_ids:
                    connection.execute(
                        "UPDATE classes SET archived_at=CURRENT_TIMESTAMP WHERE id=?",
                        (opposite_id,),
                    )
                if school.get("control_target") is not None:
                    for exam_row in connection.execute(
                        "SELECT id FROM exams WHERE grade_id=? AND archived_at IS NULL", (grade_id,)
                    ):
                        connection.execute(
                            """INSERT INTO class_targets(exam_id, class_id, line_type, target_count, source)
                               VALUES (?, ?, '特控', ?, 'base')
                               ON CONFLICT(exam_id, class_id, line_type) DO UPDATE SET
                                 target_count=excluded.target_count, source='base'""",
                            (exam_row["id"], class_id, school["control_target"]),
                        )
        for student in payload.students:
            if not all((student.get("school"), student.get("grade"), student.get("class_name"), student.get("name"), student.get("student_no"))):
                continue
            connection.execute(
                """INSERT INTO base_students(school_name, grade_name, class_name, student_name, student_no, subject_tag)
                   VALUES (?, ?, ?, ?, ?, ?)
                   ON CONFLICT(school_name, grade_name, class_name, student_no) DO UPDATE SET
                     student_name=excluded.student_name, subject_tag=excluded.subject_tag, updated_at=CURRENT_TIMESTAMP""",
                (student["school"], student["grade"], student["class_name"], student["name"], student["student_no"], student.get("subject_tag") or None),
            )
        for grade in grades:
            connection.execute(
                """UPDATE exams SET data_version=data_version+1
                   WHERE grade_id IN (SELECT id FROM grades WHERE name=?)""",
                (grade,),
            )
        connection.execute(
            "INSERT INTO operation_logs(action, target_type, summary) VALUES ('base_info_save', 'base_information', ?)",
            (f"按年级快照保存 {len(payload.schools)} 个班级、{len(payload.students)} 名学生基础信息",),
        )
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    return {"school_classes": len(payload.schools), "students": len(payload.students), "warnings": len([i for i in payload.issues if i.severity == 'warning'])}


def list_base_information(connection: sqlite3.Connection, page: int = 1, page_size: int = 12) -> dict:
    page = max(1, page)
    page_size = min(max(1, page_size), 50)
    total = connection.execute("SELECT COUNT(*) AS count FROM base_school_classes").fetchone()["count"]
    rows = connection.execute(
        """SELECT b.*, COUNT(s.id) AS student_count FROM base_school_classes b
           LEFT JOIN base_students s ON s.school_name=b.school_name AND s.grade_name=b.grade_name AND s.class_name=b.class_name
           GROUP BY b.id ORDER BY b.school_name, b.grade_name, b.class_name LIMIT ? OFFSET ?""",
        (page_size, (page - 1) * page_size),
    ).fetchall()
    items = [dict(row) for row in rows]
    ids = [item["id"] for item in items]
    targets_by_class: dict[int, list[dict]] = defaultdict(list)
    if ids:
        placeholders = ",".join("?" for _ in ids)
        for row in connection.execute(
            f"""SELECT base_class_id, target_name, target_order, target_count
                FROM base_class_targets WHERE base_class_id IN ({placeholders})
                ORDER BY target_order, target_name""",
            ids,
        ):
            targets_by_class[row["base_class_id"]].append({
                "name": row["target_name"], "order": row["target_order"], "count": row["target_count"]
            })
    for item in items:
        item["targets"] = targets_by_class.get(item["id"], [])
    return {"items": items, "page": page, "page_size": page_size, "total": total}


def build_reference_workbook_from_base_information(connection: sqlite3.Connection, output_path: str | Path) -> None:
    schools = connection.execute("SELECT * FROM base_school_classes ORDER BY school_name, grade_name, class_name").fetchall()
    students = connection.execute("SELECT * FROM base_students ORDER BY school_name, grade_name, class_name, student_no").fetchall()
    if not students:
        raise ValueError("尚未保存学生基础信息，请先上传基础信息表")
    workbook = Workbook()
    target_rows = connection.execute(
        """SELECT bct.base_class_id, bct.target_name, bct.target_order, bct.target_count
           FROM base_class_targets bct ORDER BY bct.target_order, bct.target_name"""
    ).fetchall()
    targets_by_class: dict[int, dict[str, int]] = defaultdict(dict)
    target_orders: dict[str, int] = {}
    for target in target_rows:
        targets_by_class[target["base_class_id"]][target["target_name"]] = target["target_count"]
        target_orders[target["target_name"]] = min(target_orders.get(target["target_name"], target["target_order"]), target["target_order"])
    target_names = sorted(target_orders, key=lambda name: (target_orders[name], name)) or ["特控"]
    school_sheet = workbook.active
    school_sheet.title = "学校信息"
    school_sheet.append(["省份", "学校", "年级", "班级", "方向", "班级人数", *[f"{name}目标" for name in target_names]])
    for row in schools:
        targets = targets_by_class.get(row["id"], {"特控": row["control_target"]})
        school_sheet.append([row["province"], row["school_name"], row["grade_name"], row["class_name"], row["track"], row["planned_size"], *[targets.get(name) for name in target_names]])
    student_sheet = workbook.create_sheet("学生信息")
    student_sheet.append(["学校", "年级", "班级", "学生姓名", "学号", "选科标签"])
    for row in students:
        student_sheet.append([row["school_name"], row["grade_name"], row["class_name"], row["student_name"], row["student_no"], row["subject_tag"]])
    workbook.save(output_path)


def _parse_score_sheet(
    value_workbook,
    formula_workbook,
    sheet_name: str,
    track: str,
    payload: WorkbookPayload,
) -> None:
    sheet = value_workbook[sheet_name]
    formula_sheet = formula_workbook[sheet_name]
    column_map = _build_score_column_map(formula_sheet, sheet_name, payload)
    identity_columns = column_map["identity"]
    required_identity = ["grade", "exam_name", "exam_date", "student_name"]
    required_identity.extend(["student_no", "class_name"])
    missing_identity = [IDENTITY_CANONICAL_HEADERS[key] for key in required_identity if key not in identity_columns]
    if missing_identity:
        payload.issues.append(
            ImportIssue(
                "error", "score_headers_missing",
                f"成绩表缺少可识别的身份字段：{'、'.join(missing_identity)}",
                "表头顺序可以调整，但必须保留对应含义的列名",
                sheet_name, 1,
            )
        )
        return
    for mapping in [item for item in payload.header_mappings if item["sheet"] == sheet_name]:
        payload.issues.append(
            ImportIssue(
                "warning", "header_normalized",
                f"已将非标准表头“{mapping['original']}”识别为“{mapping['normalized']}”",
                "系统已按标准字段处理，无需手工修改原文件",
                sheet_name, 1, f"第{mapping['column']}列", mapping["original"],
            )
        )
    max_column = column_map["max_column"]
    formula_rows = formula_sheet.iter_rows(min_row=2, max_col=max_column, values_only=True)

    def value_at(row: tuple[Any, ...], column: int | None) -> Any:
        return row[column - 1] if column and column <= len(row) else None

    def identity_value(row: tuple[Any, ...], field_name: str) -> Any:
        return value_at(row, identity_columns.get(field_name))

    def subject_value(row: tuple[Any, ...], code: str, field_name: str) -> Any:
        return value_at(row, column_map["subjects"][code].get(field_name))

    def total_value(row: tuple[Any, ...], field_name: str) -> Any:
        return value_at(row, column_map["totals"].get(field_name))

    score_data_columns = {
        column
        for fields in column_map["subjects"].values()
        for field_name, column in fields.items()
        if field_name in {"raw_score", "scaled_score"}
    } | {
        column for field_name, column in column_map["totals"].items()
        if field_name in {"raw_total", "scaled_total"}
    }

    for row_number, (row, formula_row) in enumerate(
        zip(sheet.iter_rows(min_row=2, max_col=max_column, values_only=True), formula_rows), start=2
    ):
        if not any(_text(value) for value in row):
            continue
        grade = _text(identity_value(row, "grade"))
        exam_name = _text(identity_value(row, "exam_name"))
        raw_exam_date = identity_value(row, "exam_date")
        exam_date = _parse_date(raw_exam_date, value_workbook.epoch)
        student_no = _student_number(identity_value(row, "student_no"))
        class_name = _text(identity_value(row, "class_name"))
        student_name = _text(identity_value(row, "student_name"))
        if not any(_text(value_at(row, column)) for column in score_data_columns):
            payload.issues.append(
                ImportIssue(
                    "warning", "score_row_without_results_skipped",
                    f"{student_name or '该学生'}本次没有任何成绩，已按未参考记录跳过",
                    "如实际参加考试，请补充至少一科成绩或总分",
                    sheet_name, row_number, "成绩",
                )
            )
            continue
        required_values = {
            "年级": grade,
            "考试名称": exam_name,
            "考试时间": exam_date,
            "姓名": student_name,
        }
        required_values.update({"学号": student_no, "班级": class_name})
        for field_name, value in required_values.items():
            if value in (None, ""):
                payload.issues.append(
                    ImportIssue(
                        "error",
                        f"score_{field_name}_missing",
                        f"成绩记录的{field_name}为空",
                        f"补充{field_name}",
                        sheet_name,
                        row_number,
                        field_name,
                        _text(raw_exam_date) if field_name == "考试时间" else None,
                    )
                )
        if _text(raw_exam_date) and exam_date is None:
            payload.issues.append(
                ImportIssue("error", "exam_date_invalid", "考试时间无法识别", "使用真实 Excel 日期或 YYYY-MM-DD", sheet_name, row_number, "考试时间", _text(raw_exam_date))
            )

        row_track = track
        subject_scores = []
        primary_code = "physics" if row_track == "物理" else "history"
        subject_codes = ["chinese", "math", "english", primary_code, "chemistry", "biology", "politics", "geography"]
        for code in subject_codes:
            fields = column_map["subjects"][code]
            raw_score = _number(subject_value(row, code, "raw_score"))
            scaled_score = _number(subject_value(row, code, "scaled_score"))
            effective_score = scaled_score if (SUBJECTS[code].category == "elective") else raw_score
            maximum = SUBJECTS[code].default_max_score
            if effective_score is not None and not 0 <= effective_score <= maximum:
                payload.issues.append(
                    ImportIssue(
                        "error",
                        "subject_score_out_of_range",
                        f"{SUBJECTS[code].name}成绩 {effective_score:g} 超出 0—{maximum:g}",
                        "核对成绩或满分配置",
                        sheet_name,
                        row_number,
                        SUBJECTS[code].name,
                        str(effective_score),
                    )
                )
            subject_scores.append(
                {
                    "subject_code": code,
                    "raw_score": raw_score,
                    "grade_level": _text(subject_value(row, code, "grade_level")) or None,
                    "scaled_score": scaled_score,
                    "effective_score": effective_score,
                    "network_rank_raw": _integer(subject_value(row, code, "network_rank_raw")),

                    "network_rank_scaled": _integer(subject_value(row, code, "network_rank_scaled")),

                }
            )
        raw_total = _number(total_value(row, "raw_total"))
        scaled_total = _number(total_value(row, "scaled_total"))
        present_scores = [item["effective_score"] for item in subject_scores if item["effective_score"] is not None]
        derived_total = round(sum(present_scores), 4) if len(present_scores) == 6 else None
        effective_total = scaled_total if scaled_total is not None else derived_total
        source = "scaled" if scaled_total is not None else "derived" if derived_total is not None else "missing"
        if effective_total is None:
            payload.issues.append(
                ImportIssue("error", "total_missing", "总分赋分及六科有效成绩合计均缺失", "补充总分赋分或完整的六科成绩及选考赋分", sheet_name, row_number, "总分")
            )
        elif not 0 <= effective_total <= 750:
            payload.issues.append(
                ImportIssue("error", "total_out_of_range", f"有效总分 {effective_total:g} 超出 0—750", "核对总分或考试满分", sheet_name, row_number, "总分", str(effective_total))
            )
        if derived_total is not None and effective_total is not None and abs(derived_total - effective_total) > 5:
            payload.issues.append(
                ImportIssue(
                    "warning",
                    "total_subject_mismatch",
                    f"有效总分与六科有效成绩合计相差 {abs(derived_total - effective_total):.1f} 分",
                    "确认是否存在赋分或模板列错位",
                    sheet_name,
                    row_number,
                    "总分",
                    str(effective_total),
                )
            )
        payload.scores.append(
            {
                "sheet_name": sheet_name,
                "row_number": row_number,
                "grade": grade,
                "exam_name": exam_name,
                "exam_date": exam_date,
                "exam_key": build_exam_key(grade, exam_name, exam_date) if grade and exam_name and exam_date else None,
                "student_no": student_no,
                "class_name": class_name,
                "student_name": student_name,
                "student_type": "普通生",
                "track": row_track,
                "raw_total": raw_total,
                "scaled_total": effective_total,
                "effective_total": effective_total,
                "total_source": source,
                "network_rank_raw": _integer(total_value(row, "network_rank_raw")),

                "network_rank_scaled": _integer(total_value(row, "network_rank_scaled")),

                "subjects": subject_scores,
            }
        )
def _parse_score_lines(workbook, payload: WorkbookPayload) -> None:
    exam_keys = {row["exam_key"] for row in payload.scores if row["exam_key"]}
    if len(exam_keys) != 1:
        if "分数线表" in workbook.sheetnames:
            payload.issues.append(
                ImportIssue(
                    "warning",
                    "reference_exam_ambiguous",
                    "分数线只能自动关联到单场考试，本文件包含零场或多场考试",
                    "建议每个文件只放一场考试，或在系统中手动补录",
                )
            )
        return
    exam_key = next(iter(exam_keys))
    if "分数线表" in workbook.sheetnames:
        sheet = workbook["分数线表"]
        row_map = {
            ("物理", "特控"): 4,
            ("物理", "211"): 5,
            ("物理", "985"): 6,
            ("历史", "特控"): 8,
            ("历史", "211"): 9,
            ("历史", "985"): 10,
        }
        for (track, line_type), row_number in row_map.items():
            value = _number(sheet.cell(row_number, 2).value)
            if value is not None:
                payload.score_lines.append(
                    {"exam_key": exam_key, "track": track, "line_type": line_type, "line_value": value}
                )
        for track in ("物理", "历史"):
            values = [entry["line_value"] for entry in payload.score_lines if entry["track"] == track]
            if values and values != sorted(values):
                payload.issues.append(
                    ImportIssue("error", "score_line_order_invalid", f"{track}方向分数线必须满足 特控 ≤ 211 ≤ 985", "核对分数线表", "分数线表")
                )
def _validate_single_school(payload: WorkbookPayload, connection=None) -> None:
    names = {_normalized_text(item.get("school")) for item in [*payload.schools, *payload.students]}
    if connection is not None:
        for query in (
            "SELECT school_name FROM base_school_classes",
            "SELECT school_name FROM base_students",
            "SELECT name FROM organizations",
        ):
            names.update(_normalized_text(row[0]) for row in connection.execute(query))
    names.discard("")
    if len(names) > 1 and not any(issue.issue_code == "single_school_required" for issue in payload.issues):
        payload.issues.append(ImportIssue(
            "error", "single_school_required", "本套系统只管理本校，检测到多个学校名称",
            "请核对基础信息中的学校名称；其他学校应独立使用一套系统"))


def _check_grade_ranks(payload: WorkbookPayload, connection, mode: str, allow_partial: bool) -> None:
    """Check optional input ranks against the resulting same-exam, same-track population."""
    scores, _ = _partition_score_rows(payload)
    grouped = defaultdict(dict)
    exam_keys = {score["exam_key"] for score in scores}
    if connection is not None and (mode != "replace" or allow_partial):
        for exam_key in exam_keys:
            rows = connection.execute(
                """SELECT er.id, se.track, er.raw_total, er.effective_total,
                          COALESCE(s.identity_key, 'legacy|' || s.student_no) AS identity_key
                   FROM exam_results er JOIN exams e ON e.id=er.exam_id
                   JOIN students s ON s.id=er.student_id
                   JOIN student_exam_enrollments se ON se.id=er.enrollment_id
                   WHERE e.exam_key=? AND er.attendance_status='normal'
                     AND er.effective_total IS NOT NULL AND se.student_type='普通生' ORDER BY er.id""",
                (exam_key,),
            ).fetchall()
            for row in rows:
                subjects = connection.execute(
                    "SELECT subject_code, raw_score, effective_score FROM subject_scores WHERE exam_result_id=?",
                    (row["id"],),
                ).fetchall()
                grouped[(exam_key, row["track"])][row["identity_key"]] = {
                    ("__TOTAL__", "raw"): row["raw_total"],
                    ("__TOTAL__", "scaled"): row["effective_total"],
                    **{(item["subject_code"], kind): item[field] for item in subjects
                       for kind, field in (("raw", "raw_score"), ("scaled", "effective_score"))},
                }
    for score in scores:
        group = grouped[(score["exam_key"], score["track"])]
        if mode != "replace" and score["identity_key"] in group:
            continue  # Append keeps the already imported student's results.
        group[score["identity_key"]] = {
            ("__TOTAL__", "raw"): score["raw_total"],
            ("__TOTAL__", "scaled"): score["effective_total"],
            **{(item["subject_code"], kind): item[field] for item in score["subjects"]
               for kind, field in (("raw", "raw_score"), ("scaled", "effective_score"))},
        }
    ranks = {}
    for key, group in grouped.items():
        fields = {field for values in group.values() for field in values}
        ranks[key] = {}
        for field in fields:
            ordered = sorted((values[field] for values in group.values()
                              if values.get(field) is not None), reverse=True)
            positions = {}
            for position, value in enumerate(ordered, 1):
                positions.setdefault(value, position)
            ranks[key][field] = positions
    for score in scores:
        key = (score["exam_key"], score["track"])
        values = grouped[key][score["identity_key"]]
        differences = []
        for item in [{"subject_code": "__TOTAL__", **score}, *score["subjects"]]:
            code = item["subject_code"]
            name = "总分" if code == "__TOTAL__" else SUBJECTS[code].name
            for kind, label in (("raw", "原始分"), ("scaled", "赋分")):
                supplied = item.get(f"network_rank_{kind}")
                computed = ranks[key].get((code, kind), {}).get(values.get((code, kind)))
                if supplied is not None and supplied != computed:
                    differences.append(f"{name}{label}：上传 {supplied}，计算 {computed if computed is not None else '无有效成绩'}")
        if differences:
            payload.issues.append(ImportIssue(
                "warning", "grade_rank_mismatch", "；".join(differences),
                "允许导入，展示以系统计算的年级排名为准；预览按追加范围校验，覆盖时重新校验",
                score["sheet_name"], score["row_number"], "年级排名"))


def parse_workbook(path: str | Path, connection: sqlite3.Connection | None = None, *, mode: str = "append", allow_partial: bool = False) -> WorkbookPayload:
    started = perf_counter()
    formula_workbook = load_workbook(path, read_only=True, data_only=False)
    value_workbook = load_workbook(path, read_only=True, data_only=True)
    payload = WorkbookPayload(sheets_found=list(formula_workbook.sheetnames))
    parse_started = perf_counter()
    _parse_school_sheet(value_workbook, payload)
    _parse_student_sheet(value_workbook, payload)
    for sheet_name, track in SCORE_SHEETS.items():
        if sheet_name not in formula_workbook.sheetnames:
            payload.issues.append(
                ImportIssue("warning", "score_sheet_missing", f"缺少{sheet_name}，该方向不会导入", sheet_name=sheet_name)
            )
            continue
        _parse_score_sheet(value_workbook, formula_workbook, sheet_name, track, payload)
    payload.timings["数据解析"] = round(perf_counter() - parse_started, 3)
    lines_started = perf_counter()
    _parse_score_lines(value_workbook, payload)
    payload.timings["附加数据解析"] = round(perf_counter() - lines_started, 3)
    formula_workbook.close()
    value_workbook.close()
    validation_started = perf_counter()
    _validate_single_school(payload, connection)
    _cross_validate(payload, connection)
    _check_grade_ranks(payload, connection, mode, allow_partial)
    payload.timings["身份与数据质检"] = round(perf_counter() - validation_started, 3)
    payload.timings["预检解析合计"] = round(perf_counter() - started, 3)
    return payload


def _candidate_parts(item: dict) -> tuple[str, str, str]:
    return (
        _student_number(item.get("student_no")),
        _normalized_text(item.get("class_name")).casefold(),
        _normalized_text(item.get("name")).casefold(),
    )


def _candidate_indexes(records: list[dict]) -> dict[str, dict]:
    indexes: dict[str, dict] = {"number_name": {}, "name_class": {}}

    def append_unique(index: dict, key: tuple, item: dict) -> None:
        bucket = index.setdefault(key, [])
        marker = item.get("identity_key") or (
            *_candidate_parts(item),
            _normalized_text(item.get("school")).casefold(),
        )
        if not any(
            (existing.get("identity_key") or (
                *_candidate_parts(existing),
                _normalized_text(existing.get("school")).casefold(),
            )) == marker
            for existing in bucket
        ):
            bucket.append(item)

    for item in records:
        number, class_name, name = _candidate_parts(item)
        if number and name:
            append_unique(indexes["number_name"], (number.casefold(), name), item)
        if name and class_name:
            append_unique(indexes["name_class"], (name, class_name), item)
    return indexes


def _existing_candidate_records(connection: sqlite3.Connection | None, scores: list[dict]) -> list[dict]:
    if connection is None:
        return []
    student_numbers = sorted({_student_number(score.get("student_no")) for score in scores if _student_number(score.get("student_no"))})
    if not student_numbers:
        return []
    output: list[dict] = []
    for start in range(0, len(student_numbers), 900):
        values = student_numbers[start:start + 900]
        placeholders = ",".join("?" for _ in values)
        output.extend(
            dict(row) for row in connection.execute(
                f"""
                SELECT COALESCE(s.display_student_no, s.student_no) AS student_no,
                       s.id AS student_id, s.identity_key, s.name, s.student_type,
                       o.name AS school, g.name AS grade, c.class_name
                FROM students s
                LEFT JOIN student_exam_enrollments see ON see.student_id=s.id
                LEFT JOIN organizations o ON o.id=see.organization_id
                LEFT JOIN grades g ON g.id=see.grade_id
                LEFT JOIN classes c ON c.id=see.class_id
                WHERE s.archived_at IS NULL
                  AND COALESCE(s.display_student_no, s.student_no) IN ({placeholders})
                GROUP BY s.id, o.name, g.name, c.class_name
                """,
                values,
            ).fetchall()
        )
    return output


def _cross_validate(payload: WorkbookPayload, connection: sqlite3.Connection | None) -> None:
    masters = [student for student in payload.students if student.get("student_no")]
    master_indexes = _candidate_indexes(masters)
    existing_indexes = _candidate_indexes(_existing_candidate_records(connection, payload.scores))
    seen_exam_student: set[tuple[str, str]] = set()

    def candidates_for(score: dict) -> tuple[list[dict], str]:
        number = _student_number(score.get("student_no"))
        class_name = _normalized_text(score.get("class_name")).casefold()
        name = _normalized_text(score.get("student_name")).casefold()
        if number and name:
            key = (number.casefold(), name)
            exact = master_indexes["number_name"].get(key, [])
            if not exact:
                exact = existing_indexes["number_name"].get(key, [])
            return exact, "number_name" if len(exact) == 1 else "ambiguous" if exact else "not_found"
        # Only a genuinely blank/formula-missing student number may use the current
        # base-information fallback. A non-empty but wrong number is never guessed.
        name_class = master_indexes["name_class"].get((name, class_name), [])
        return (
            name_class,
            "name_class_fallback" if len(name_class) == 1 else "ambiguous" if name_class else "not_found",
        )

    for score in payload.scores:
        candidates, matched_by = candidates_for(score)
        if not candidates:
            payload.issues.append(
                ImportIssue(
                    "error",
                    "student_not_in_master",
                    "成绩记录无法在学生信息中定位到同一学生",
                    "核对学号和姓名；永久身份按学号＋姓名识别",
                    score["sheet_name"],
                    score["row_number"],
                    "学号",
                    score.get("student_no"),
                )
            )
            continue
        if len(candidates) > 1:
            payload.issues.append(
                ImportIssue(
                    "error", "student_identity_ambiguous",
                    "当前成绩行无法按学号和姓名唯一定位学生",
                    "核对重复的学号＋姓名；学号为空时再核对姓名＋班级",
                    score["sheet_name"], score["row_number"], "学号", score.get("student_no"),
                )
            )
            continue
        master = candidates[0]
        identity_key = master.get("identity_key") or _regular_identity_key(
            master.get("student_no"), master.get("name")
        )
        master["identity_key"] = identity_key
        score["identity_key"] = identity_key
        score["master"] = master
        if not master.get("subject_tag"):
            # 选科标签不是上传前提。依据实际有分数的首选/再选科目补齐，
            # 例如物理、化学、生物有成绩即记录为“物化生”。
            tag_map = {"physics": "物", "history": "史", "chemistry": "化", "biology": "生", "politics": "政", "geography": "地"}
            master["subject_tag"] = "".join(
                tag_map[item["subject_code"]]
                for item in score["subjects"]
                if item["subject_code"] in tag_map and item["effective_score"] is not None
            )
        selected = set(_subject_codes_from_tag(master.get("subject_tag", ""), score["track"]))
        selected.update(item["subject_code"] for item in score["subjects"]
                        if (SUBJECTS[item["subject_code"]].category == "elective") and
                        (item["raw_score"] is not None or item["scaled_score"] is not None))
        for subject in score["subjects"]:
            code = subject["subject_code"]
            if code in selected and (SUBJECTS[code].category == "elective") and subject["scaled_score"] is None:
                payload.issues.append(ImportIssue(
                    "error", "selected_scaled_score_missing",
                    f"{SUBJECTS[code].name}为已选科目，但赋分缺失",
                    "补齐赋分后重新导入；系统不会用原始分替代", score["sheet_name"],
                    score["row_number"], f"{SUBJECTS[code].name}赋分"))
        if not score.get("student_no"):
            score["student_no"] = _student_number(master.get("student_no"))
            payload.issues = [
                issue for issue in payload.issues
                if not (
                    issue.issue_code == "score_学号_missing"
                    and issue.sheet_name == score["sheet_name"]
                    and issue.row_number == score["row_number"]
                )
            ]
        if matched_by == "name_class_fallback":
            payload.issues.append(
                ImportIssue(
                    "warning", "student_identity_fallback",
                    "学号公式缓存缺失，已通过本次基础信息中唯一的姓名＋班级定位学生",
                    "建议在 Excel 中重新计算并保存公式，提交前核对本行",
                    score["sheet_name"], score["row_number"], "学号", score.get("student_no"),
                )
            )
        master_name = master.get("name")
        if master_name and score["student_name"] and _normalized_text(master_name) != _normalized_text(score["student_name"]):
            payload.issues.append(
                ImportIssue(
                    "warning",
                    "student_name_mismatch",
                    f"同一复合身份的姓名不一致：学生信息为“{master_name}”，成绩表为“{score['student_name']}”",
                    "请核对学号、学校、班级和姓名",
                    score["sheet_name"],
                    score["row_number"],
                    "姓名",
                    score["student_name"],
                )
            )
        key = (score["exam_key"], identity_key)
        if all(key) and key in seen_exam_student:
            payload.issues.append(
                ImportIssue(
                    "error",
                    "duplicate_exam_student",
                    "同一考试中同一复合身份出现重复成绩",
                    "删除重复的学号＋姓名成绩行",
                    score["sheet_name"],
                    score["row_number"],
                    "学生身份",
                    identity_key,
                )
            )
        seen_exam_student.add(key)

    # Perform the same-exam number/name conflict check after blank formula IDs
    # have been recovered, so partial import cannot admit only one side of a conflict.
    exam_number_names: dict[tuple[str, str], set[str]] = {}
    for score in payload.scores:
        if not score.get("master"):
            continue
        number = _student_number(score.get("student_no"))
        name = _normalized_text(score.get("student_name")).casefold()
        if number and name:
            exam_number_names.setdefault(
                (score.get("exam_key", ""), number.casefold()), set()
            ).add(name)
    conflicting_numbers = {
        key for key, names in exam_number_names.items() if len(names) > 1
    }
    for score in payload.scores:
        key = (
            score.get("exam_key", ""),
            _student_number(score.get("student_no")).casefold(),
        )
        if score.get("master") and key in conflicting_numbers:
            payload.issues.append(
                ImportIssue(
                    "error", "exam_student_number_name_conflict",
                    "同一场考试中同一学号对应了多个不同姓名",
                    "核对本场考试中的学号和姓名；系统不会静默合并或覆盖",
                    score["sheet_name"], score["row_number"], "学号", score.get("student_no"),
                )
            )


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


SCORE_SHEET_LEVEL_ERROR_CODES = {
    "score_header_duplicate",
    "score_headers_missing",
    }


def _partition_score_rows(payload: WorkbookPayload) -> tuple[list[dict], list[dict]]:
    """Split parsed score rows into rows that are safe to persist and blocked rows."""
    score_sheets = set(SCORE_SHEETS)
    if any(issue.issue_code == "single_school_required" for issue in payload.issues):
        return [], list(payload.scores)
    blocked_sheets = {
        issue.sheet_name
        for issue in payload.issues
        if issue.severity == "error"
        and issue.sheet_name in score_sheets
        and (
            issue.issue_code in SCORE_SHEET_LEVEL_ERROR_CODES
            or issue.row_number in {None, 1}
        )
    }
    blocked_rows = {
        (issue.sheet_name, issue.row_number)
        for issue in payload.issues
        if issue.severity == "error"
        and issue.sheet_name in score_sheets
        and issue.row_number is not None
        and issue.row_number > 1
    }
    valid: list[dict] = []
    blocked: list[dict] = []
    for score in payload.scores:
        location = (score.get("sheet_name"), score.get("row_number"))
        master = score.get("master") or {}
        structurally_valid = all(
            [
                score.get("exam_key"),
                score.get("identity_key"),
                master.get("name"),
                score.get("grade"),
                score.get("exam_name"),
                score.get("exam_date"),
                score.get("student_name"),
                score.get("track") in TRACKS,
                score.get("effective_total") is not None,
            ]
        )
        if (
            score.get("sheet_name") in blocked_sheets
            or location in blocked_rows
            or not structurally_valid
        ):
            blocked.append(score)
        else:
            valid.append(score)
    return valid, blocked


def preview_workbook(
    connection: sqlite3.Connection,
    source_path: str | Path,
    original_filename: str,
) -> dict:
    started = perf_counter()
    path = Path(source_path)
    payload = parse_workbook(path, connection)
    file_hash = _file_sha256(path)
    duplicate = connection.execute(
        "SELECT id, batch_key, committed_at FROM import_batches WHERE source_sha256=? AND status LIKE 'committed%' LIMIT 1",
        (file_hash,),
    ).fetchone()
    if duplicate:
        payload.issues.append(
            ImportIssue(
                "warning",
                "duplicate_file",
                "检测到相同文件曾经成功导入",
                "提交时选择覆盖、追加或取消",
            )
        )
    errors = [issue for issue in payload.issues if issue.severity == "error"]
    warnings = [issue for issue in payload.issues if issue.severity == "warning"]
    valid_scores, blocked_scores = _partition_score_rows(payload)
    batch_key = uuid.uuid4().hex
    summary = {
        "sheets": payload.sheets_found,
        "school_rows": len(payload.schools),
        "student_rows": len(payload.students),
        "score_rows": len(payload.scores),
        "valid_score_rows": len(valid_scores),
        "blocked_score_rows": len(blocked_scores),
        "score_line_rows": len(payload.score_lines),
        "header_mappings": payload.header_mappings,
        "duplicate_committed_batch": duplicate["batch_key"] if duplicate else None,
        "timings": dict(payload.timings),
    }
    database_started = perf_counter()
    cursor = connection.execute(
        """
        INSERT INTO import_batches(
          batch_key, original_filename, source_path, source_sha256, source_size,
          status, row_count, valid_row_count, error_count, warning_count, summary_json
        ) VALUES (?, ?, ?, ?, ?, 'previewed', ?, ?, ?, ?, ?)
        """,
        (
            batch_key,
            original_filename,
            str(path),
            file_hash,
            path.stat().st_size,
            len(payload.scores),
            len(valid_scores),
            len(errors),
            len(warnings),
            json.dumps(summary, ensure_ascii=False),
        ),
    )
    batch_id = cursor.lastrowid
    connection.executemany(
        """
        INSERT INTO import_issues(
          import_batch_id, severity, sheet_name, row_number, field_name,
          raw_value, issue_code, message, suggestion
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (
                batch_id,
                issue.severity,
                issue.sheet_name,
                issue.row_number,
                issue.field_name,
                issue.raw_value,
                issue.issue_code,
                issue.message,
                issue.suggestion,
            )
            for issue in payload.issues
        ],
    )
    connection.commit()
    summary["timings"]["预检记录保存"] = round(perf_counter() - database_started, 3)
    summary["timings"]["预检合计"] = round(perf_counter() - started, 3)
    connection.execute(
        "UPDATE import_batches SET summary_json=? WHERE id=?",
        (json.dumps(summary, ensure_ascii=False), batch_id),
    )
    connection.commit()
    LOGGER.info(
        "import_preview rows=%s valid=%s blocked=%s errors=%s warnings=%s elapsed=%.3fs",
        len(payload.scores), len(valid_scores), len(blocked_scores), len(errors), len(warnings),
        summary["timings"]["预检合计"],
    )
    return {
        "batch_key": batch_key,
        "original_filename": original_filename,
        "status": "blocked" if errors else "ready",
        "summary": summary,
        "error_count": len(errors),
        "warning_count": len(warnings),
        "valid_row_count": len(valid_scores),
        "blocked_row_count": len(blocked_scores),
        "issues": [asdict(issue) for issue in payload.issues[:200]],
        "issues_truncated": len(payload.issues) > 200,
    }


def attach_batch_preparation(
    connection: sqlite3.Connection,
    batch_key: str,
    preparation: dict,
    source_files: list[str],
) -> dict:
    batch = connection.execute(
        "SELECT summary_json FROM import_batches WHERE batch_key=?", (batch_key,)
    ).fetchone()
    if batch is None:
        raise ValueError("导入批次不存在")
    summary = json.loads(batch["summary_json"] or "{}")
    summary["preparation"] = preparation
    summary["timings"] = {
        **(preparation.get("timings") or {}),
        **(summary.get("timings") or {}),
    }
    summary["timings"]["导入预检总耗时"] = round(sum(
        value for key, value in summary["timings"].items()
        if key in {"规范表生成合计", "预检合计"} and isinstance(value, (int, float))
    ), 3)
    summary["source_files"] = source_files
    connection.execute(
        "UPDATE import_batches SET summary_json=? WHERE batch_key=?",
        (json.dumps(summary, ensure_ascii=False), batch_key),
    )
    connection.commit()
    preview = get_batch_preview(connection, batch_key)
    preview["status"] = "blocked" if preview["error_count"] else "ready"
    return preview


def _get_or_create_id(connection: sqlite3.Connection, table: str, key: dict, defaults: dict | None = None) -> int:
    values = {**key, **(defaults or {})}
    columns = ",".join(values)
    placeholders = ",".join("?" for _ in values)
    connection.execute(
        f"INSERT OR IGNORE INTO {table} ({columns}) VALUES ({placeholders})",
        tuple(values.values()),
    )
    where = " AND ".join(f"{column}=?" for column in key)
    row = connection.execute(
        f"SELECT id FROM {table} WHERE {where}", tuple(key.values())
    ).fetchone()
    return int(row["id"])


def commit_batch(
    connection: sqlite3.Connection,
    batch_key: str,
    mode: str = "reject",
    allow_partial: bool = False,
    warnings_confirmed: bool | None = None,
) -> dict:
    if mode not in {"reject", "append", "replace"}:
        raise ValueError("mode 必须是 reject、append 或 replace")
    batch = connection.execute(
        "SELECT * FROM import_batches WHERE batch_key=?", (batch_key,)
    ).fetchone()
    if batch is None:
        raise ValueError("导入批次不存在")
    if str(batch["status"]).startswith("committed"):
        raise ValueError("该批次已经提交")
    if batch["warning_count"] and warnings_confirmed is False:
        raise ValueError("请先确认允许入库但需核对的提醒")
    if batch["error_count"] and not allow_partial:
        raise ValueError("导入预检存在阻止入库的问题")
    payload = parse_workbook(batch["source_path"], connection, mode=mode, allow_partial=allow_partial)
    current_errors = [issue for issue in payload.issues if issue.severity == "error"]
    if current_errors and not allow_partial:
        raise ValueError("重新校验时发现阻止入库的问题")
    valid_scores, blocked_scores = _partition_score_rows(payload)
    if allow_partial:
        if not valid_scores:
            raise ValueError("当前批次没有可安全保存的有效成绩行")
        payload.scores = valid_scores
        if any(issue.issue_code == "score_line_order_invalid" for issue in current_errors):
            payload.score_lines = []
    partial_commit = bool(allow_partial and (batch["error_count"] or current_errors))
    blocked_row_count = len(blocked_scores) if allow_partial else 0
    committed = 0
    skipped = 0
    try:
        connection.execute("BEGIN IMMEDIATE")
        exam_ids: dict[str, int] = {}
        # “覆盖已有成绩”应替换整场考试快照。旧实现只覆盖身份键完全相同的行，
        # 学生改号或基础信息纠正后会留下旧成绩并再插入新成绩，造成历史人数膨胀。
        if mode == "replace" and not allow_partial:
            for exam_key in {score["exam_key"] for score in payload.scores if score.get("exam_key")}:
                existing_exam = connection.execute(
                    "SELECT id FROM exams WHERE exam_key=?", (exam_key,)
                ).fetchone()
                if existing_exam is None:
                    continue
                exam_id = existing_exam["id"]
                connection.execute("DELETE FROM exam_results WHERE exam_id=?", (exam_id,))
                connection.execute("DELETE FROM student_exam_enrollments WHERE exam_id=?", (exam_id,))
                exam_ids[exam_key] = exam_id
        for score in payload.scores:
            master = score.get("master")
            if master is None or not score.get("identity_key"):
                raise ValueError(f"第 {score['row_number']} 行缺少已确认的学生复合身份")
            organization_name = master.get("school") or "未配置学校"
            class_name = master.get("class_name") or score["class_name"]
            organization_id = _get_or_create_id(
                connection,
                "organizations",
                {"name": organization_name},
                {"province": None},
            )
            grade_id = _get_or_create_id(connection, "grades", {"name": score["grade"]})
            class_id = _get_or_create_id(
                connection,
                "classes",
                {
                    "organization_id": organization_id,
                    "grade_id": grade_id,
                    "class_name": class_name,
                    "track": score["track"],
                },
            )
            student_id = _get_or_create_id(
                connection,
                "students",
                {"student_no": score["identity_key"]},
                {
                    "display_student_no": score.get("student_no") or None,
                    "identity_key": score["identity_key"],
                    "student_type": master.get("student_type") or score.get("student_type") or "普通生",
                    "name": master["name"],
                },
            )
            connection.execute(
                """
                UPDATE students
                SET name=?, display_student_no=?, identity_key=?, student_type=?
                WHERE id=?
                """,
                (
                    master["name"], score.get("student_no") or None, score["identity_key"],
                    master.get("student_type") or score.get("student_type") or "普通生", student_id,
                ),
            )
            if score["exam_key"] not in exam_ids:
                connection.execute(
                    """
                    INSERT OR IGNORE INTO exams(
                      exam_key, grade_id, name, exam_date, total_max_score
                    ) VALUES (?, ?, ?, ?, 750)
                    """,
                    (
                        score["exam_key"],
                        grade_id,
                        score["exam_name"],
                        score["exam_date"].isoformat(),
                    ),
                )
                exam_ids[score["exam_key"]] = connection.execute(
                    "SELECT id FROM exams WHERE exam_key=?", (score["exam_key"],)
                ).fetchone()["id"]
            exam_id = exam_ids[score["exam_key"]]
            existing = connection.execute(
                "SELECT id FROM exam_results WHERE exam_id=? AND student_id=?",
                (exam_id, student_id),
            ).fetchone()
            if existing and mode == "reject":
                raise ValueError("数据库中已有同一考试同一复合身份；请选择覆盖、追加或取消")
            if existing and mode == "append":
                skipped += 1
                continue
            if existing and mode == "replace":
                connection.execute("DELETE FROM exam_results WHERE id=?", (existing["id"],))
            elective_codes = _subject_codes_from_tag(master.get("subject_tag", ""), score["track"])
            connection.execute(
                """
                INSERT INTO student_exam_enrollments(
                  student_id, exam_id, organization_id, grade_id, class_id, track,
                  elective_subjects_json, student_type, source
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'import')
                ON CONFLICT(student_id, exam_id) DO UPDATE SET
                  organization_id=excluded.organization_id,
                  grade_id=excluded.grade_id,
                  class_id=excluded.class_id,
                  track=excluded.track,
                  elective_subjects_json=excluded.elective_subjects_json,
                  student_type=excluded.student_type,
                  source='import'
                """,
                (
                    student_id,
                    exam_id,
                    organization_id,
                    grade_id,
                    class_id,
                    score["track"],
                    json.dumps(elective_codes, ensure_ascii=False),
                    master.get("student_type") or score.get("student_type") or "普通生",
                ),
            )
            enrollment_id = connection.execute(
                "SELECT id FROM student_exam_enrollments WHERE student_id=? AND exam_id=?",
                (student_id, exam_id),
            ).fetchone()["id"]
            cursor = connection.execute(
                """
                INSERT INTO exam_results(
                  exam_id, student_id, enrollment_id, raw_total, scaled_total,
                  effective_total, total_source, network_rank_raw,
                  network_rank_scaled, attendance_status,
                  import_batch_id
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'normal', ?)
                """,
                (
                    exam_id,
                    student_id,
                    enrollment_id,
                    score["raw_total"],
                    score["scaled_total"],
                    score["effective_total"],
                    score["total_source"],
                    score["network_rank_raw"],

                    score["network_rank_scaled"],

                    batch["id"],
                ),
            )
            result_id = cursor.lastrowid
            for subject in score["subjects"]:
                if subject["effective_score"] is None:
                    continue
                connection.execute(
                    """
                    INSERT INTO subject_scores(
                      exam_result_id, subject_code, raw_score, grade_level, scaled_score,
                      effective_score, network_rank_raw,
                      network_rank_scaled
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        result_id,
                        subject["subject_code"],
                        subject["raw_score"],
                        subject["grade_level"],
                        subject["scaled_score"],
                        subject["effective_score"],
                        subject["network_rank_raw"],

                        subject["network_rank_scaled"],

                    ),
                )
            committed += 1
        for line in payload.score_lines:
            exam_id = exam_ids.get(line["exam_key"])
            if exam_id:
                connection.execute(
                    """
                    INSERT INTO score_lines(
                      exam_id, track, line_type, subject_code, line_value, source, rule_version
                    ) VALUES (?, ?, ?, '__TOTAL__', ?, 'import', ?)
                    ON CONFLICT(exam_id, track, line_type, subject_code)
                    DO UPDATE SET line_value=excluded.line_value, source='import', rule_version=excluded.rule_version
                    """,
                    (exam_id, line["track"], line["line_type"], line["line_value"], CORE_RULE_VERSION),
                )
        for school in payload.schools:
            if not all(
                [
                    school.get("school"),
                    school.get("grade"),
                    school.get("class_name"),
                ]
            ):
                continue
            organization_id = _get_or_create_id(
                connection, "organizations", {"name": school["school"]}, {"province": None}
            )
            grade_id = _get_or_create_id(connection, "grades", {"name": school["grade"]})
            tracks = {school["track"]} if school.get("track") in TRACKS else {
                score["track"] for score in payload.scores
                if score.get("grade") == school["grade"] and score.get("master", {}).get("school") == school["school"]
                and score.get("master", {}).get("class_name") == school["class_name"]
            }
            for school_track in tracks:
                class_id = _get_or_create_id(
                    connection, "classes",
                    {"organization_id": organization_id, "grade_id": grade_id,
                     "class_name": school["class_name"], "track": school_track},
                )
                connection.execute(
                    "UPDATE classes SET planned_size=COALESCE(?, planned_size), control_target=COALESCE(?, control_target) WHERE id=?",
                    (school["planned_size"], school["control_target"], class_id),
                )
                for target in _school_targets(school):
                    connection.execute(
                        """INSERT INTO class_stage_targets(
                             class_id, target_name, target_order, target_count, source
                           ) VALUES (?, ?, ?, ?, 'import')
                           ON CONFLICT(class_id, target_name) DO UPDATE SET
                             target_order=excluded.target_order,
                             target_count=excluded.target_count,
                             source='import', updated_at=CURRENT_TIMESTAMP""",
                        (class_id, target["name"], target["order"], target["count"]),
                    )
                if school["control_target"] is None:
                    continue
                matching_exam_ids = {
                    exam_ids[score["exam_key"]] for score in payload.scores
                    if score["grade"] == school["grade"] and score["track"] == school_track and score["exam_key"] in exam_ids
                }
                for exam_id in matching_exam_ids:
                    connection.execute(
                        """INSERT INTO class_targets(exam_id, class_id, line_type, target_count, source)
                           VALUES (?, ?, '特控', ?, 'import')
                           ON CONFLICT(exam_id, class_id, line_type)
                           DO UPDATE SET target_count=excluded.target_count, source='import'""",
                        (exam_id, class_id, school["control_target"]),
                    )
        for exam_id in set(exam_ids.values()):
            connection.execute("UPDATE exams SET data_version=data_version+1 WHERE id=?", (exam_id,))
        target_grades = sorted({school["grade"] for school in payload.schools if _school_targets(school)})
        for grade_name in target_grades:
            connection.execute(
                """UPDATE exams SET data_version=data_version+1
                   WHERE grade_id IN (SELECT id FROM grades WHERE name=?)
                     AND id NOT IN ({})""".format(
                         ",".join("?" for _ in set(exam_ids.values())) or "NULL"
                     ),
                [grade_name, *set(exam_ids.values())],
            )
        final_status = "committed_partial" if partial_commit else "committed"
        # Replacement can change the ranking population after preview. Preserve the final check.
        connection.execute("DELETE FROM import_issues WHERE import_batch_id=? AND issue_code='grade_rank_mismatch'", (batch["id"],))
        rank_warnings = [issue for issue in payload.issues if issue.issue_code == "grade_rank_mismatch"]
        connection.executemany(
            """INSERT INTO import_issues(import_batch_id, severity, sheet_name, row_number,
               field_name, issue_code, message, suggestion) VALUES (?, 'warning', ?, ?, ?, ?, ?, ?)""",
            [(batch["id"], issue.sheet_name, issue.row_number, issue.field_name, issue.issue_code,
              issue.message, "按最终导入范围校验；允许入库，展示以系统计算的年级排名为准") for issue in rank_warnings],
        )
        connection.execute(
            "UPDATE import_batches SET warning_count=(SELECT COUNT(*) FROM import_issues WHERE import_batch_id=? AND severity='warning') WHERE id=?",
            (batch["id"], batch["id"]),
        )
        connection.execute(
            """
            UPDATE import_batches
            SET status=?, committed_at=CURRENT_TIMESTAMP, valid_row_count=?
            WHERE id=?
            """,
            (final_status, committed, batch["id"]),
        )
        connection.execute(
            "INSERT INTO operation_logs(action, target_type, target_id, summary) VALUES ('import_commit', 'import_batch', ?, ?)",
            (
                batch_key,
                f"提交 {committed} 条成绩，重复跳过 {skipped} 条，阻止项跳过 {blocked_row_count} 条",
            ),
        )
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    return {
        "batch_key": batch_key,
        "status": "committed_partial" if partial_commit else "committed",
        "committed_rows": committed,
        "skipped_rows": skipped,
        "blocked_rows_skipped": blocked_row_count,
        "rank_warning_count": len(rank_warnings),
        "exam_keys": list(dict.fromkeys(score["exam_key"] for score in payload.scores)),
    }


def get_batch_preview(connection: sqlite3.Connection, batch_key: str) -> dict:
    batch = connection.execute(
        """
        SELECT batch_key, original_filename, status, row_count, valid_row_count,
               error_count, warning_count, summary_json, created_at, committed_at
        FROM import_batches WHERE batch_key=?
        """,
        (batch_key,),
    ).fetchone()
    if batch is None:
        raise ValueError("导入批次不存在")
    summary = json.loads(batch["summary_json"] or "{}")
    valid_row_count = summary.get("valid_score_rows", batch["valid_row_count"])
    blocked_row_count = summary.get(
        "blocked_score_rows", max(0, batch["row_count"] - valid_row_count)
    )
    return {
        "batch_key": batch["batch_key"],
        "original_filename": batch["original_filename"],
        "status": batch["status"],
        "summary": summary,
        "error_count": batch["error_count"],
        "warning_count": batch["warning_count"],
        "valid_row_count": valid_row_count,
        "blocked_row_count": blocked_row_count,
        "created_at": batch["created_at"],
        "committed_at": batch["committed_at"],
    }


def list_batch_issues(
    connection: sqlite3.Connection,
    batch_key: str,
    severity: str,
    page: int = 1,
    page_size: int = 50,
) -> dict:
    if severity not in {"error", "warning"}:
        raise ValueError("severity 必须是 error 或 warning")
    batch = connection.execute(
        "SELECT id FROM import_batches WHERE batch_key=?", (batch_key,)
    ).fetchone()
    if batch is None:
        raise ValueError("导入批次不存在")
    page = max(1, page)
    page_size = max(1, min(page_size, 100))
    total = connection.execute(
        "SELECT COUNT(*) AS count FROM import_issues WHERE import_batch_id=? AND severity=?",
        (batch["id"], severity),
    ).fetchone()["count"]
    total_pages = max(1, math.ceil(total / page_size))
    page = min(page, total_pages)
    rows = connection.execute(
        """
        SELECT severity, sheet_name, row_number, field_name, raw_value,
               issue_code, message, suggestion
        FROM import_issues
        WHERE import_batch_id=? AND severity=?
        ORDER BY COALESCE(sheet_name, ''), COALESCE(row_number, 0), id
        LIMIT ? OFFSET ?
        """,
        (batch["id"], severity, page_size, (page - 1) * page_size),
    ).fetchall()
    return {
        "items": [dict(row) for row in rows],
        "severity": severity,
        "page": page,
        "page_size": page_size,
        "total": total,
        "total_pages": total_pages,
    }


def list_batches(connection: sqlite3.Connection, limit: int = 100) -> list[dict]:
    rows = connection.execute(
        """
        SELECT batch_key, original_filename, source_sha256, source_size, status,
               row_count, valid_row_count, error_count, warning_count,
               created_at, committed_at, summary_json
        FROM import_batches
        WHERE status LIKE 'committed%'
        ORDER BY id DESC LIMIT ?
        """,
        (max(1, min(limit, 100)),),
    ).fetchall()
    output = []
    for row in rows:
        item = dict(row)
        item["summary"] = json.loads(item.pop("summary_json") or "{}")
        output.append(item)
    return output


def _batch_file_paths(row: sqlite3.Row) -> list[str]:
    paths = [row["source_path"]]
    try:
        summary = json.loads(row["summary_json"] or "{}")
        paths.extend(summary.get("source_files") or [])
    except (TypeError, ValueError, json.JSONDecodeError):
        pass
    return list(dict.fromkeys(str(path) for path in paths if path))


def _cleanup_orphan_import_data(connection: sqlite3.Connection) -> None:
    connection.execute(
        """
        DELETE FROM student_exam_enrollments
        WHERE NOT EXISTS(
          SELECT 1 FROM exam_results er
          WHERE er.enrollment_id=student_exam_enrollments.id
        )
        """
    )
    connection.execute(
        """
        DELETE FROM exams
        WHERE NOT EXISTS(SELECT 1 FROM exam_results er WHERE er.exam_id=exams.id)
        """
    )
    connection.execute(
        """
        DELETE FROM students
        WHERE NOT EXISTS(
          SELECT 1 FROM student_exam_enrollments se WHERE se.student_id=students.id
        )
        AND NOT EXISTS(SELECT 1 FROM exam_results er WHERE er.student_id=students.id)
        """
    )
    connection.execute(
        """
        DELETE FROM classes
        WHERE NOT EXISTS(
          SELECT 1 FROM student_exam_enrollments se WHERE se.class_id=classes.id
        )
        AND NOT EXISTS(SELECT 1 FROM class_targets ct WHERE ct.class_id=classes.id)
        AND NOT EXISTS(SELECT 1 FROM class_stage_targets cst WHERE cst.class_id=classes.id)
        """
    )


def delete_import_batch(connection: sqlite3.Connection, batch_key: str) -> dict:
    row = connection.execute(
        "SELECT * FROM import_batches WHERE batch_key=?", (batch_key,)
    ).fetchone()
    if row is None:
        raise ValueError("导入批次不存在")
    file_paths = _batch_file_paths(row)
    result_count = connection.execute(
        "SELECT COUNT(*) AS count FROM exam_results WHERE import_batch_id=?",
        (row["id"],),
    ).fetchone()["count"]
    try:
        connection.execute("BEGIN IMMEDIATE")
        affected_exams = [
            item["exam_id"] for item in connection.execute(
                "SELECT DISTINCT exam_id FROM exam_results WHERE import_batch_id=?",
                (row["id"],),
            )
        ]
        connection.execute(
            "DELETE FROM exam_results WHERE import_batch_id=?", (row["id"],)
        )
        if affected_exams:
            placeholders = ",".join("?" for _ in affected_exams)
            connection.execute(
                f"UPDATE exams SET data_version=data_version+1 WHERE id IN ({placeholders})",
                affected_exams,
            )
        connection.execute("DELETE FROM import_batches WHERE id=?", (row["id"],))
        _cleanup_orphan_import_data(connection)
        connection.execute(
            "INSERT INTO operation_logs(action, target_type, target_id, summary) VALUES ('import_delete', 'import_batch', ?, ?)",
            (batch_key, f"删除导入批次及其当前关联的 {result_count} 条成绩"),
        )
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    return {
        "batch_key": batch_key,
        "deleted_results": result_count,
        "file_paths": file_paths,
    }


def clear_import_cache(connection: sqlite3.Connection) -> dict:
    batches = connection.execute("SELECT * FROM import_batches ORDER BY id").fetchall()
    file_paths = list(
        dict.fromkeys(path for row in batches for path in _batch_file_paths(row))
    )
    result_count = connection.execute(
        "SELECT COUNT(*) AS count FROM exam_results WHERE import_batch_id IS NOT NULL"
    ).fetchone()["count"]
    try:
        connection.execute("BEGIN IMMEDIATE")
        connection.execute("DELETE FROM exam_results WHERE import_batch_id IS NOT NULL")
        connection.execute("DELETE FROM import_batches")
        _cleanup_orphan_import_data(connection)
        connection.execute(
            "INSERT INTO operation_logs(action, target_type, target_id, summary) VALUES ('import_cache_clear', 'import_batch', 'all', ?)",
            (f"清除 {len(batches)} 个上传批次及 {result_count} 条当前关联成绩",),
        )
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    return {
        "deleted_batches": len(batches),
        "deleted_results": result_count,
        "file_paths": file_paths,
    }
